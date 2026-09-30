#!/usr/bin/env python3
"""Vista #42438 cigarette count 2026-09-29 — S2K site 293 post + Jayden packet."""
from __future__ import annotations

import json
import shutil
import sys
import urllib.parse
from datetime import datetime
from decimal import Decimal, ROUND_HALF_UP
from pathlib import Path

from fpdf import FPDF
from openpyxl import load_workbook

sys.path.insert(0, "/tmp/cig-inv")
import build_all  # noqa: E402

WORKDIR = Path("/tmp/cig-inv/vista-count")
OUT = Path("/workspace/vista-cigarette-count/42438/2026-09-29")
STATION = "42438"
NAME = "Vista"
SITEID = 293
ACCOUNT = -121
TRANREF = "CIG COUNT 09292026"
TRANDATE = "2026-09-29"
REPORT_END = "2026-09-29"

# From CIG SHEET 09292026.pdf (4 pages + margin note).
HANDWRITTEN: dict[int, int] = {
    1: 5,
    2: 8,
    3: 10,
    4: 9,
    5: 17,
    6: 9,
    7: 22,
    8: 3,
    9: 13,
    10: 8,
    11: 18,
    12: 12,
    13: 2,
    14: 0,
    15: 19,
    16: 0,
    17: 22,
    18: 23,
    19: 18,
    20: 12,
    21: 13,
    22: 16,
    23: 9,
    24: 10,
    25: 8,
    26: 4,
    27: 8,
    28: 4,
    29: 8,
    30: 10,
    31: 10,
    32: 4,
    33: 13,
    34: 14,
    35: 20,
    36: 15,
    37: 2,
    38: 27,
    39: 6,
    40: 14,
    41: 10,
    42: 16,
    43: 19,
    44: 8,
    45: 35,
    46: 8,
    47: 12,
    48: 15,
    49: 5,
    50: 7,
    51: 10,
    52: 7,
    53: 28,
    54: 0,
    55: 10,
    56: 16,
    57: 11,
    58: 0,
    59: 0,
    60: 41,
    61: 27,
    62: 2,
    63: 15,
    64: 22,
    65: 6,
    66: 39,
    67: 21,
    68: 40,
    69: 3,
    70: 8,
    71: 3,
    72: 12,
    73: 6,
    74: 10,
    75: 13,
    76: 11,
    77: 6,
    78: 10,
    79: 5,
    80: 10,
    81: 3,
    82: 9,
    83: 36,
    84: 2,
    85: 8,
    86: 4,
    87: 20,
    88: 23,
    89: 13,
    90: 13,
    91: 9,
    92: 5,
    93: 6,
    94: 30,
    95: 9,
    96: 20,
    97: 5,
    98: 27,
    99: 9,
    100: 10,
    101: 15,
    102: 12,
    103: 5,
    104: 3,
    105: 8,
    106: 10,
    107: 16,
    108: 10,
    109: 11,
    110: 3,
    111: 3,
    112: 34,
    113: 30,
    114: 15,
}

# Margin (not on printed 114-line sheet).
MARGIN_EXTRA = [
    {
        "upc": "047995200404",
        "marginNoteUpc": "047995205404",
        "desc": "American Spirit sky",
        "qty": 8,
    }
]

BELOW_GRID_SKIPPED: list[str] = []


def merged_counts() -> dict[int, int]:
    return dict(HANDWRITTEN)


def money(n: float) -> str:
    sign = "-" if n < 0 else ""
    return f"{sign}${abs(n):,.2f}"


def qty(n) -> str:
    if isinstance(n, float) and n.is_integer():
        return str(int(n))
    return str(n)


def round2(x: float) -> float:
    return float(Decimal(str(x)).quantize(Decimal("0.01"), rounding=ROUND_HALF_UP))


def load_sheet_items() -> list[dict]:
    rows = []
    for line in (WORKDIR / "sheet-items.tsv").read_text().splitlines():
        n, upc, desc = line.split("\t", 2)
        rows.append({"n": int(n), "upc": upc.strip(), "desc": desc.strip()})
    return rows


def pull_packs_xlsx(session) -> Path:
    build_all.END = REPORT_END
    depts = build_all.cig_depts(session)
    pack_dept = next((d for d, n in depts if "PACK" in n.upper()), depts[0][0])
    sid = build_all.session_id(session)
    url = (
        "https://store.s2kprime.com/report?rpt="
        + urllib.parse.quote("SKU Sales By Item")
        + "&id="
        + urllib.parse.quote(str(sid))
        + f"&StartDate=2026-08-01&EndDate={REPORT_END}&StationID={SITEID}&Dept={pack_dept}&Toggle=1&rs:Format=EXCELOPENXML"
    )
    r = session.get(url, timeout=240)
    if r.content[:2] != b"PK":
        raise RuntimeError(f"packs report failed {r.status_code} site {SITEID}")
    path = WORKDIR / "packs.xlsx"
    path.write_bytes(r.content)
    return path


def parse_qoh(path: Path) -> dict[str, dict]:
    wb = load_workbook(path, data_only=True)
    ws = wb.active
    latest: dict[str, dict] = {}
    for row in ws.iter_rows(min_row=13, values_only=True):
        if not row or len(row) < 9:
            continue
        desc, sku, dept, sdate, qoh, sold, cost = (
            row[1],
            row[3],
            row[4],
            row[5],
            row[6],
            row[7],
            row[8],
        )
        if not isinstance(sku, str) or not sku.strip().isdigit():
            continue
        if not isinstance(sdate, datetime):
            continue
        upc = sku.strip()
        rec = {
            "desc": (desc or "").strip(),
            "dept": (dept or "").strip(),
            "dt": sdate.date().isoformat(),
            "qoh": float(qoh or 0),
            "cost": float(cost or 0),
        }
        prev = latest.get(upc)
        if not prev or rec["dt"] >= prev["dt"]:
            latest[upc] = rec
    return latest


def product_lookup(session, upc: str) -> dict | None:
    r = session.get(
        "https://store.s2kprime.com/api/product",
        params={"upc": upc},
        timeout=40,
    )
    if r.status_code != 200:
        return None
    data = r.json()
    if not isinstance(data, list) or not data:
        return None
    for hit in data:
        if str(hit.get("upc") or "") == upc:
            return hit
    return data[0]


def find_existing_tran(session, tranref: str = TRANREF) -> dict | None:
    xsrf = session.cookies.get("XSRF-TOKEN")
    hdr = {"X-XSRF-TOKEN": xsrf} if xsrf else {}
    for url in (
        f"https://store.s2kprime.com/api/tran/item-10?siteid={SITEID}&fromdate={TRANDATE}&todate={TRANDATE}",
        f"https://store.s2kprime.com/api/tran/item?siteid={SITEID}&trandate={TRANDATE}",
    ):
        r = session.get(url, headers=hdr, timeout=60)
        if r.status_code != 200:
            continue
        try:
            data = r.json()
        except json.JSONDecodeError:
            continue
        items = data if isinstance(data, list) else data.get("data") or []
        for it in items:
            if str(it.get("tranref") or "") == tranref and int(it.get("siteid") or 0) == SITEID:
                return it
    return None


def post_inventory(session, lines: list[dict]) -> dict:
    meta_path = WORKDIR / "post-meta.json"
    if meta_path.exists():
        try:
            prior = json.loads(meta_path.read_text())
            if (
                prior.get("tranref") == TRANREF
                and int(prior.get("siteid") or 0) == SITEID
                and prior.get("id")
            ):
                return {"skipped": True, "id": prior["id"], "reason": "already posted (local post-meta)"}
        except (json.JSONDecodeError, TypeError, ValueError):
            pass
    existing = find_existing_tran(session, TRANREF)
    if existing:
        return {"skipped": True, "id": existing.get("id"), "reason": "already posted"}

    tran_lines = []
    for row in lines:
        tran_lines.append(
            {
                "varid": row["varid"],
                "qty": float(row["actual"]),
                "cal_type": 0,
                "props": json.dumps({"sold": 0}),
                "model_dirty": True,
            }
        )
    item = {
        "siteid": SITEID,
        "tranref": TRANREF,
        "trandate": TRANDATE,
        "trandate2": TRANDATE,
        "trantype": 10,
        "lines": tran_lines,
        "deleted_lines": [],
        "props": json.dumps(
            {
                "note": (
                    f"{NAME} cigarette physical count written {TRANDATE}. "
                    f"Station {STATION}, S2K site {SITEID}. "
                    "Quantity on hand set to counted packs from CIG SHEET 09292026."
                )
            }
        ),
        "model_dirty": True,
    }
    body = {
        "type": "item",
        "trantype": 10,
        "data": [item],
        "create_new": True,
        "return_get": True,
    }
    xsrf = session.cookies.get("XSRF-TOKEN")
    hdr = {"X-XSRF-TOKEN": xsrf} if xsrf else {}
    r = session.post(
        "https://store.s2kprime.com/api/tran/update/item-10-undefined",
        json=body,
        headers=hdr,
        timeout=120,
    )
    if r.status_code == 200 and "error" not in r.text.lower():
        data = r.json()
        if isinstance(data, list) and data:
            return {"skipped": False, "id": data[0].get("id"), "response": data[0]}
    raise RuntimeError(f"post failed {r.status_code} {r.text[:400]}")


class Sheet(FPDF):
    def footer(self):
        self.set_y(-12)
        self.set_font("Helvetica", "", 8)
        self.set_x(self.l_margin)
        self.cell(self.epw, 8, f"{self.page_no()}/{{nb}}", align="R")


def write_pdfs(lines: list[dict], summary: dict, tranid: int | None):
    cols = [12, 86, 16, 16, 14, 20, 26]
    headers = ["#", "Description", "Actual", "Book", "Diff", "Unit", "Amount"]

    def table(pdf, rows, amount_key):
        pdf.set_font("Helvetica", "B", 8)
        pdf.set_x(pdf.l_margin)
        for w, label in zip(cols, headers):
            pdf.cell(w, 5, label, new_x="RIGHT", new_y="TOP")
        pdf.ln(5)
        pdf.set_font("Helvetica", "", 8)
        for row in rows:
            if pdf.get_y() > 262:
                pdf.add_page()
            vals = [
                str(row["n"]),
                row["desc"][:42],
                qty(row["actual"]),
                qty(row["qoh"]),
                qty(row["diff"]),
                f"{row['cost']:.4f}",
                money(row[amount_key]),
            ]
            pdf.set_x(pdf.l_margin)
            for w, val in zip(cols, vals):
                pdf.cell(w, 4.4, val, new_x="RIGHT", new_y="TOP")
            pdf.ln(4.4)

    pdf = Sheet(format="Letter", unit="mm")
    pdf.alias_nb_pages()
    pdf.set_auto_page_break(auto=True, margin=14)
    pdf.set_margins(10, 12, 10)
    pdf.add_page()
    pdf.set_font("Helvetica", "B", 14)
    pdf.cell(0, 8, f"{NAME} #{STATION} cigarette count (S2K site {SITEID})", new_x="LMARGIN", new_y="NEXT")
    pdf.set_font("Helvetica", "", 10)
    pdf.multi_cell(
        0,
        5,
        f"September 29, 2026. Actual {money(summary['actualCost'])} for {summary['actualPacks']} packs. "
        f"Overall {money(summary['overall'])}. Book QoH from S2K site {SITEID} packs report through {REPORT_END}.",
    )
    actual_rows = sorted(lines, key=lambda r: (0 if isinstance(r["n"], int) else 1, r["n"]))
    for row in actual_rows:
        row["actual_amt"] = round2(row["actual"] * row["cost"])
    pdf.ln(2)
    pdf.set_font("Helvetica", "B", 11)
    pdf.cell(0, 6, "Actual cost", new_x="LMARGIN", new_y="NEXT")
    table(pdf, actual_rows, "actual_amt")
    minus_rows = sorted([r for r in lines if r["cost_delta"] < 0], key=lambda r: r["cost_delta"])
    plus_rows = sorted([r for r in lines if r["cost_delta"] > 0], key=lambda r: -r["cost_delta"])
    pdf.ln(2)
    pdf.set_font("Helvetica", "B", 11)
    pdf.cell(0, 6, "Minus, book above the count", new_x="LMARGIN", new_y="NEXT")
    table(pdf, minus_rows, "cost_delta")
    pdf.ln(2)
    pdf.set_font("Helvetica", "B", 11)
    pdf.cell(0, 6, "Plus, count above the book", new_x="LMARGIN", new_y="NEXT")
    table(pdf, plus_rows, "cost_delta")
    pdf.output(OUT / "count.pdf")

    pdf2 = Sheet(format="Letter", unit="mm")
    pdf2.alias_nb_pages()
    pdf2.set_auto_page_break(auto=True, margin=14)
    pdf2.set_margins(10, 12, 10)
    pdf2.add_page()
    pdf2.set_font("Helvetica", "B", 14)
    pdf2.cell(0, 8, "SKU Inventory", new_x="LMARGIN", new_y="NEXT")
    pdf2.set_font("Helvetica", "B", 12)
    pdf2.cell(0, 6, f"{NAME} #{STATION}", new_x="LMARGIN", new_y="NEXT")
    pdf2.set_font("Helvetica", "", 10)
    pdf2.cell(0, 5, f"September 29, 2026. Reference {TRANREF}. Site {SITEID}.", new_x="LMARGIN", new_y="NEXT")
    if tranid:
        pdf2.cell(0, 5, f"Transaction {tranid}.", new_x="LMARGIN", new_y="NEXT")
    cols2 = [12, 100, 36, 36]
    pdf2.ln(2)
    pdf2.set_font("Helvetica", "B", 8)
    for w, label in zip(cols2, ["#", "Description", "UPC", "Quantity on hand"]):
        pdf2.cell(w, 5, label, new_x="RIGHT", new_y="TOP")
    pdf2.ln(5)
    pdf2.set_font("Helvetica", "", 8)
    for row in sorted(lines, key=lambda r: (0 if isinstance(r["n"], int) else 1, r["n"])):
        if pdf2.get_y() > 262:
            pdf2.add_page()
        pdf2.set_x(pdf2.l_margin)
        for w, val in zip(
            cols2,
            [str(row["n"]), row["desc"][:48], row["upc"], qty(row["actual"])],
        ):
            pdf2.cell(w, 4.4, val, new_x="RIGHT", new_y="TOP")
        pdf2.ln(4.4)
    pdf2.output(OUT / "s2k-change.pdf")


def build_entry(summary: dict, tranid: int | None) -> dict:
    base = f"/inventory/{STATION}/{TRANDATE}"
    return {
        "stationId": STATION,
        "client": NAME,
        "date": TRANDATE,
        "dateLabel": "September 29, 2026",
        "category": "Cigarettes",
        "heading": (
            f"Count written September 29, 2026. SKU Inventory {TRANREF}, S2K site {SITEID}"
            + (f", transaction {tranid}." if tranid else ".")
            + " Physical count from CIG SHEET 09292026 compared to Vista S2K QoH."
        ),
        "summary": summary["narrative"],
        "count": {
            "actualPacks": summary["actualPacks"],
            "actualCost": summary["actualCost"],
            "minusPacks": summary["minusPacks"],
            "minusCost": summary["minusCost"],
            "plusPacks": summary["plusPacks"],
            "plusCost": summary["plusCost"],
            "overall": summary["overall"],
        },
        "files": [
            {
                "group": "Scans",
                "label": "Manager count sheet (CIG SHEET 09292026)",
                "href": f"{base}/count-sheet.pdf",
            },
            {
                "group": "S2K quantity on hand",
                "label": "Packs, department 200",
                "href": f"{base}/qoh-packs.pdf",
            },
            {
                "group": "S2K quantity on hand",
                "label": "Cartons, department 281",
                "href": f"{base}/qoh-cartons.pdf",
            },
            {
                "group": "S2K quantity on hand",
                "label": "Quantity on hand change",
                "href": f"{base}/s2k-change.pdf",
            },
            {
                "group": "Count",
                "label": "Cigarette count",
                "href": f"{base}/count.pdf",
            },
        ],
    }


def main():
    import argparse

    ap = argparse.ArgumentParser()
    ap.add_argument("--no-post", action="store_true")
    args = ap.parse_args()
    counts = merged_counts()
    OUT.mkdir(parents=True, exist_ok=True)
    items = load_sheet_items()
    if len(items) != 114:
        raise SystemExit(f"expected 114 sheet rows, got {len(items)}")
    logins = build_all.load_logins()
    session, _ = build_all.login(*logins["hotmail"])
    build_all.switch(session, ACCOUNT)

    packs_path = pull_packs_xlsx(session)
    qoh_by_upc = parse_qoh(packs_path)
    (WORKDIR / "latest-qoh.json").write_text(json.dumps(qoh_by_upc, indent=2))

    lines: list[dict] = []
    skipped_no_varid = []
    for it in items:
        n = it["n"]
        actual = counts.get(n)
        if actual is None:
            continue
        prod = product_lookup(session, it["upc"])
        if not prod:
            skipped_no_varid.append(n)
            continue
        q = qoh_by_upc.get(it["upc"], {})
        qoh = float(q.get("qoh", 0))
        cost = float(q.get("cost") or 0)
        if cost <= 0 and prod.get("cost"):
            cost = float(prod["cost"])
        diff = float(actual) - qoh
        lines.append(
            {
                "n": n,
                "upc": it["upc"],
                "desc": it["desc"],
                "actual": actual,
                "qoh": qoh,
                "diff": diff,
                "cost": cost,
                "cost_delta": round2(diff * cost),
                "varid": prod["id"],
                "qoh_asof": q.get("dt"),
                "dept": q.get("dept", "CIGARETTES - PACKS"),
            }
        )

    actual_packs = sum(r["actual"] for r in lines)
    actual_cost = round2(sum(r["actual"] * r["cost"] for r in lines))
    minus = [r for r in lines if r["diff"] < 0]
    plus = [r for r in lines if r["diff"] > 0]
    minus_packs = sum(-r["diff"] for r in minus)
    plus_packs = sum(r["diff"] for r in plus)
    minus_cost = round2(sum(-r["cost_delta"] for r in minus if r["cost_delta"] < 0))
    plus_cost = round2(sum(r["cost_delta"] for r in plus if r["cost_delta"] > 0))
    overall = round2(plus_cost - minus_cost)

    if args.no_post:
        post = {"skipped": True, "reason": "--no-post"}
        tranid = None
        meta_path = WORKDIR / "post-meta.json"
        if meta_path.exists():
            try:
                tranid = json.loads(meta_path.read_text()).get("id")
            except json.JSONDecodeError:
                pass
    else:
        post = post_inventory(session, lines)
        tranid = post.get("id")

    summary = {
        "station": STATION,
        "name": NAME,
        "siteid": SITEID,
        "date": TRANDATE,
        "tranref": TRANREF,
        "tranid": tranid,
        "linesPosted": len(lines),
        "marginSkipped": MARGIN_EXTRA,
        "belowGridSkipped": BELOW_GRID_SKIPPED,
        "actualPacks": int(actual_packs),
        "actualCost": actual_cost,
        "minusPacks": int(minus_packs),
        "minusCost": minus_cost,
        "plusPacks": int(plus_packs),
        "plusCost": plus_cost,
        "overall": overall,
        "post": post,
        "narrative": (
            f"Vista S2K site {SITEID}: actual {money(actual_cost)} for {int(actual_packs)} packs on "
            f"{len(lines)} lines. Minus {money(minus_cost)} ({int(minus_packs)} packs), "
            f"plus {money(plus_cost)} ({int(plus_packs)} packs). Overall {money(overall)}."
        ),
    }
    (WORKDIR / "summary.json").write_text(json.dumps(summary, indent=2))
    (OUT / "lines.json").write_text(json.dumps(lines, indent=2))
    if not args.no_post and tranid:
        (WORKDIR / "post-meta.json").write_text(
            json.dumps(
                {
                    "id": tranid,
                    "tranref": TRANREF,
                    "trandate": TRANDATE,
                    "siteid": SITEID,
                    "nlines": len(lines),
                    "post": post,
                },
                indent=2,
            )
        )

    write_pdfs(lines, summary, tranid)

    sid = build_all.session_id(session)
    for dept_id, fname in (("200", "qoh-packs.pdf"), ("281", "qoh-cartons.pdf")):
        url = (
            "https://store.s2kprime.com/report?rpt="
            + urllib.parse.quote("SKU Sales By Item")
            + "&id="
            + urllib.parse.quote(str(sid))
            + f"&StartDate=2026-08-01&EndDate={REPORT_END}&StationID={SITEID}&Dept={dept_id}&Toggle=1&rs:Format=PDF"
        )
        rp = session.get(url, timeout=240)
        if rp.content[:4] == b"%PDF":
            (OUT / fname).write_bytes(rp.content)

    shutil.copy(WORKDIR / "CIG_SHEET_09292026.pdf", OUT / "count-sheet.pdf")
    (OUT / "entry.json").write_text(json.dumps(build_entry(summary, tranid), indent=2) + "\n")

    readme = Path("/workspace/vista-cigarette-count/README.md")
    readme.write_text(
        f"""# {NAME} cigarette count — September 29, 2026

Station {STATION}. **S2K site {SITEID}** (Vista only — not Koval/Spring).

- Reference: `{TRANREF}`
- Transaction: {tranid or '(see post-meta.json)'}

## Totals (vs Vista S2K QoH)

- Actual: {summary['actualPacks']} packs, {money(summary['actualCost'])}
- Minus: {summary['minusPacks']} packs, {money(summary['minusCost'])}
- Plus: {summary['plusPacks']} packs, {money(summary['plusCost'])}
- Overall: {money(summary['overall'])}

Packet: `42438/2026-09-29/`.
"""
    )

    print(json.dumps(summary, indent=2))
    if skipped_no_varid:
        print("no varid for rows", skipped_no_varid, file=sys.stderr)


if __name__ == "__main__":
    main()
