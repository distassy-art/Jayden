#!/usr/bin/env python3
"""Fill Daily Excel from SoftSP: Daily Book sales + DLY/DPT purchases (Deduct rules).

**Always run after SoftSP pulls** (DailyTotal+Summary + DLY + DPT):

  SoftSP PDFs (daily + dly + dpt)
    → this script (sales + Net Daily Purchases via Deduct)
    → * Daily.xlsx (OneDrive)  ← purchases always overwritten
    → sync_from_excel.py --publish
    → website MTD

Purchases (always updated when ``--dly-dir`` is set):
  - Pull both ``*dly.pdf`` (dated invoice expand) and ``*dpt.pdf``
    (vendor/dept rollup) every SoftSP cycle — same folder tree.
  - Net Daily Purchases (column F) = Deduct rules applied to **DLY**
    invoice lines (``scripts/store_purchase_rules.py`` / Excel Deduct sheet).
  - DPT stays on disk for dept/vendor audit and Financial Audit leave-2;
    it does not replace dated DLY nets for F.
  - Empty countable day → blank F (not 0, not leftover formula).
  - Sales cells stay fill-once; **purchases always refresh**.

Usage:
  python3 scripts/fill_daily_dly_dpt.py --kinds dly,dpt
  python3 scripts/fill_daily_excel_from_softsp.py \\
    --xlsx-dir /tmp/books_margin_fresh \\
    --pdf-dir /tmp/s2k/exports/daily_0923 \\
    --dly-dir /tmp/s2k/exports/dly_dpt_run \\
    --days 2026-09-22,2026-09-23 \\
    --out-dir /tmp/books_fill_0923 \\
    --refresh-rules

  python3 scripts/sync_from_excel.py --xlsx-dir /tmp/books_fill_0923 \\
    --month 2026-09 --publish
"""
from __future__ import annotations

import argparse
import json
import re
import shutil
import sys
from datetime import date, datetime
from pathlib import Path

try:
    import openpyxl
    from openpyxl.utils import get_column_letter
except ImportError as exc:  # pragma: no cover
    raise SystemExit("openpyxl required: pip install openpyxl") from exc

try:
    import pdfplumber
except ImportError as exc:  # pragma: no cover
    raise SystemExit("pdfplumber required: pip install pdfplumber") from exc

REPO = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(REPO / "scripts"))
from store_purchase_rules import (  # noqa: E402
    apply_invoices,
    load_deduct_from_workbook,
    load_snapshot,
    rules_to_json,
    scan_excel_dir,
)

HDR = 8  # September sheet header row; day 1 → row 9

# id → (name, xlsx filename, pdf_mode, daily_subdir, tso, dly_subdir)
STORES = {
    "42004": ("Arco Placentia", "Arco Placentia Daily.xlsx", "single", "42004_Pl", None, "42004_Pl"),
    "42179": ("Arco HB", "Arco HB Daily.xlsx", "single", "42179_HB", None, "42179_HB"),
    "42352": ("Arco Db", "Arco Db Daily.xlsx", "single", "42352_Db", None, "42352_Db"),
    "42674": ("Tustin", "Tustin Daily.xlsx", "single", "42674_Tu", None, "42674_Tu"),
    "42021": ("Westminster", "Westminster Daily.xlsx", "bd", "BD_central", "42021", "BD_central"),
    "42048": ("San Diego", "San Diego Daily.xlsx", "bd", "BD_central", "42048", "BD_central"),
    "42098": ("Brookhurst 75", "Brookhurst 75 Daily.xlsx", "bd", "BD_central", "42098", "BD_central"),
    "42279": ("Koval", "Koval Daily.xlsx", "bd", "BD_central", "42279", "BD_central"),
    "42280": ("Spring Mtn", "Spring Mtn Daily.xlsx", "bd", "BD_central", "42280", "BD_central"),
    "42281": ("Charleston", "Charleston Daily.xlsx", "bd", "BD_central", "42281", "BD_central"),
    "42282": ("Oakey Las Vegas Blvd", "Oakey Las Vegas Blvd Daily.xlsx", "bd", "BD_central", "42282", "BD_central"),
    "42399": ("Garden Grove", "Garden Grove Daily.xlsx", "bd", "BD_central", "42399", "BD_central"),
    "42438": ("Vista", "Vista Daily.xlsx", "bd", "BD_central", "42438", "BD_central"),
    "42439": ("Lamb", "Lamb Daily.xlsx", "bd", "BD_central", "42439", "BD_central"),
}

TSO_LINE = re.compile(
    r"TSO #(\d+)\s+#\d+\s+\S+\s+(\d{2}/\d{2}/\d{2})\s+(-?\$[\d,]+\.\d{2})"
)


def money(s: str) -> float | None:
    s = (s or "").strip().replace(",", "").replace("$", "")
    s = s.replace("(", "-").replace(")", "")
    if s in ("", "-", "--"):
        return None
    try:
        return float(s)
    except ValueError:
        return None


def parse_dept_amount(text: str, label: str) -> float:
    m = re.search(rf"{re.escape(label)}\s+[\d,.]+\s+\$([\d,.]+)", text, re.I)
    return (money(m.group(1)) or 0.0) if m else 0.0


def parse_cstore_fields(text: str) -> dict:
    out: dict = {}
    m = re.search(
        r"Station Total:\s*([\d,.]+)\s+\$([\d,.]+)\s+([\d.]+)%\s+\$([\-\d,.]+)",
        text,
    )
    if m:
        out["cstore_total"] = money(m.group(2))
    out["tax1"] = parse_dept_amount(text, "TAX 1")
    out["tax4"] = parse_dept_amount(text, "TAX 4")
    out["scratch"] = parse_dept_amount(text, "SCRATCH TICKETS")
    out["lotto"] = parse_dept_amount(text, "LOTTO")
    card = parse_dept_amount(text, "CARD ACTIVATIONS")
    if card == 0.0:
        card = parse_dept_amount(text, "CARD ACTIVATION")
    out["card"] = card
    return out


def parse_fuel(text: str) -> dict:
    out: dict = {}
    m = re.search(
        r"Grand Total:\s*([\d,.]+)\s+\$([\d,.]+)\s+\$([\d,.]+)\s+\$([\-\d,.]+)",
        text,
    )
    if not m:
        m = re.search(
            r"Station Total:\s*([\d,.]+)\s+\$([\d,.]+)\s+\$([\-\d,.]+)\s+\$([\-\d,.]+)",
            text,
        )
    if m:
        out["gas_vol"] = float(m.group(1).replace(",", ""))
        out["gas_profit"] = money(m.group(4))
    return out


def parse_single(path: Path) -> dict:
    with pdfplumber.open(path) as pdf:
        text = "\n".join(p.extract_text() or "" for p in pdf.pages)
    fuel_part = text
    for marker in ("CStore Sales", "C-Store Sales"):
        if marker in text:
            fuel_part = text.split(marker, 1)[0]
            break
    out = parse_fuel(fuel_part)
    cstore = text
    for marker in ("CStore Sales", "C-Store Sales", "Department Qty"):
        if marker in text:
            cstore = text.split(marker, 1)[1]
            break
    out.update(parse_cstore_fields(cstore))
    return out


def parse_bd_store(path: Path, tso: str) -> dict | None:
    with pdfplumber.open(path) as pdf:
        text = "\n".join(p.extract_text() or "" for p in pdf.pages)
    chunks = re.split(r"(?=TSO #\d+)", text)
    fuel = None
    cstore = None
    for part in chunks:
        m = re.match(r"TSO #(\d+)", part)
        if not m:
            continue
        raw = m.group(1)
        if not (raw == tso or raw.startswith(tso)):
            continue
        if re.search(
            r"Station Total:\s*[\d,.]+\s+\$[\d,.]+\s+[\d.]+%\s+\$", part
        ):
            cstore = part
        elif "Unleaded" in part or "Diesel" in part:
            fuel = part
        elif "TAX 1" in part or "SCRATCH" in part or "CANDY" in part:
            cstore = part
        elif fuel is None:
            fuel = part
    if fuel is None and cstore is None:
        return None
    out: dict = {}
    if fuel:
        out.update(parse_fuel(fuel))
    if cstore:
        out.update(parse_cstore_fields(cstore))
    return out


def parse_dly_invoices(path: Path, store_ids: set[str] | None = None) -> list[dict]:
    """Parse SoftSP None-Fuel Invoice Total Expand PDF → invoice lines.

    Each line: ``{date, vendor, amount, store_id}``.
    """
    with pdfplumber.open(path) as pdf:
        text = "\n".join(p.extract_text() or "" for p in pdf.pages)
    invoices: list[dict] = []
    vendor: str | None = None
    for line in text.splitlines():
        line = line.strip()
        if not line:
            continue
        if line.startswith("TSO #"):
            m = TSO_LINE.search(line)
            if not m:
                continue
            sid, ds, amt = m.group(1), m.group(2), m.group(3)
            if store_ids and sid not in store_ids and not any(
                sid.startswith(s) for s in store_ids
            ):
                continue
            mm, dd, yy = ds.split("/")
            iso = f"{2000 + int(yy):04d}-{int(mm):02d}-{int(dd):02d}"
            amount = float(amt.replace("$", "").replace(",", ""))
            invoices.append(
                {
                    "store_id": sid,
                    "date": iso,
                    "vendor": vendor or "UNKNOWN",
                    "amount": amount,
                }
            )
            continue
        # Vendor header: NAME $total $… (skip report chrome)
        low = line.lower()
        if any(
            x in low
            for x in (
                "grouped by",
                "inv #",
                "inv date",
                "station(s)",
                "last updated",
                "of ",
                "non-fuel",
                "page ",
                "report total",
                "grand total",
                "station total",
            )
        ):
            continue
        if re.search(r"\$[\d,]+\.\d{2}", line) and re.search(r"[A-Za-z]", line):
            name = re.split(r"\s+-?\$", line, 1)[0].strip()
            if name and len(name) < 70 and not name.upper().startswith("TSO"):
                vendor = re.sub(r"\s+", " ", name)
    return invoices


def pdf_stamp(d: date) -> str:
    return f"{d.month:02d}{d.day:02d}{d.year}.pdf"


def dly_stamp(d: date) -> str:
    return f"{d.month:02d}{d.day:02d}{d.year}dly.pdf"


def header_map(ws) -> dict[str, str]:
    cols: dict[str, str] = {}
    for c in range(1, 16):
        h = ws.cell(HDR, c).value
        if not isinstance(h, str):
            continue
        hl = h.lower()
        letter = get_column_letter(c)
        if "gas volume" in hl:
            cols["gas_vol"] = letter
        elif "gas profit" in hl:
            cols["gas_profit"] = letter
        elif "net daily purchases" in hl or hl.strip().startswith("net purchases"):
            cols["purch"] = letter
        elif "c-store total" in hl or "reported c-store total" in hl:
            cols["cstore_total"] = letter
        elif hl.strip().startswith("tax 1"):
            cols["tax1"] = letter
        elif hl.strip().startswith("tax 4"):
            cols["tax4"] = letter
        elif "scratch" in hl:
            cols["scratch"] = letter
        elif "lotto" in hl:
            cols["lotto"] = letter
        elif "card" in hl:
            cols["card"] = letter
    return cols


def remap_formula(formula: str, src_row: int, dst_row: int) -> str:
    return re.sub(rf"(?<![A-Z0-9]){src_row}(?!\d)", str(dst_row), formula)


def template_formulas(ws, day: int) -> dict[str, str]:
    target = HDR + day
    for ref_day in range(1, 32):
        rr = HDR + ref_day
        dval = ws[f"D{rr}"].value
        if ws[f"B{rr}"].value in (None, ""):
            continue
        if not (isinstance(dval, str) and dval.startswith("=")):
            continue
        mapping = {}
        for col in ("D", "E", "G", "H", "I", "L"):
            val = ws[f"{col}{rr}"].value
            if isinstance(val, str) and val.startswith("="):
                mapping[col] = remap_formula(val, rr, target)
        return mapping
    return {}


def write_values(
    ws, day: int, vals: dict, cols: dict[str, str], purch: float | None = None,
    *, update_purch: bool = False,
) -> list[str]:
    r = HDR + day
    changed = []
    for field, key in [
        ("gas_vol", "gas_vol"),
        ("gas_profit", "gas_profit"),
        ("cstore_total", "cstore_total"),
        ("tax1", "tax1"),
        ("tax4", "tax4"),
        ("scratch", "scratch"),
        ("lotto", "lotto"),
        ("card", "card"),
    ]:
        if field not in cols:
            continue
        val = vals.get(key)
        if val is None:
            continue
        cell = f"{cols[field]}{r}"
        if ws[cell].value not in (None, ""):
            continue
        ws[cell] = val
        changed.append(f"{cell}={val}")
    # Purchases: always overwrite when update_purch (None → blank, not leftover formula)
    purch_col = cols.get("purch") or "F"
    if update_purch or purch is not None:
        cell = f"{purch_col}{r}"
        ws[cell] = purch
        changed.append(f"{cell}=purch:{purch}")
    for col, formula in template_formulas(ws, day).items():
        if col == "F":
            continue
        if ws[f"{col}{r}"].value in (None, ""):
            ws[f"{col}{r}"] = formula
            changed.append(f"{col}=formula")
    return changed


def fill_brookhurst(
    wb, day: int, vals: dict, purch: float | None = None, *, update_purch: bool = False
) -> list[str]:
    src = wb["September 2026 Source"]
    ws = wb["September 2026"]
    sr = 2 + day
    mr = HDR + day
    changed = []
    if src[f"A{sr}"].value in (None, ""):
        src[f"A{sr}"] = datetime(2026, 9, day)
        changed.append(f"src!A{sr}=date")
    tax_sum = sum(
        (vals.get(k) or 0.0) for k in ("tax1", "tax4", "scratch", "lotto", "card")
    )
    mapping = {
        "B": vals.get("gas_vol"),
        "C": vals.get("gas_profit"),
        "D": (
            round(vals["gas_profit"] / vals["gas_vol"], 4)
            if vals.get("gas_vol") and vals.get("gas_profit") is not None
            else None
        ),
        "E": (
            round(vals["cstore_total"] - tax_sum, 2)
            if vals.get("cstore_total") is not None
            else None
        ),
        "J": vals.get("cstore_total"),
        "K": vals.get("tax1", 0.0),
        "L": vals.get("tax4", 0.0),
        "M": vals.get("scratch", 0.0),
        "N": vals.get("lotto", 0.0),
        "O": vals.get("card", 0.0),
    }
    for col, val in mapping.items():
        if val is None:
            continue
        if src[f"{col}{sr}"].value in (None, ""):
            src[f"{col}{sr}"] = val
            changed.append(f"src!{col}{sr}={val}")
    if update_purch or purch is not None:
        src[f"F{sr}"] = purch
        changed.append(f"src!F{sr}=purch:{purch}")
        ws[f"F{mr}"] = f"='September 2026 Source'!F{sr}"
        changed.append(f"F{mr}->src purch")
    for col, formula in {
        "G": f'=IF(COUNTA(A{sr}:F{sr})=0,"",E{sr}-F{sr})',
        "H": f'=IFERROR(G{sr}/E{sr},"")',
        "I": f'=IF(COUNTA(A{sr}:F{sr})=0,"",C{sr}+G{sr})',
    }.items():
        if src[f"{col}{sr}"].value in (None, ""):
            src[f"{col}{sr}"] = formula
    links = {
        "B": f"='September 2026 Source'!B{sr}",
        "C": f"='September 2026 Source'!C{sr}",
        "E": f"='September 2026 Source'!E{sr}",
        "J": f"='September 2026 Source'!J{sr}",
        "K": f"='September 2026 Source'!K{sr}",
        "L": f"='September 2026 Source'!L{sr}",
        "M": f"='September 2026 Source'!M{sr}",
        "N": f"='September 2026 Source'!N{sr}",
        "O": f"='September 2026 Source'!O{sr}",
    }
    for col, formula in links.items():
        if ws[f"{col}{mr}"].value in (None, ""):
            ws[f"{col}{mr}"] = formula
            changed.append(f"{col}{mr}->src")
    for col, formula in template_formulas(ws, day).items():
        if col in ("D", "G", "H", "I") and ws[f"{col}{mr}"].value in (None, ""):
            ws[f"{col}{mr}"] = formula
            changed.append(f"{col}=formula")
    return changed


def fill_garden_grove(
    ws, day: int, vals: dict, purch: float | None = None, *, update_purch: bool = False
) -> list[str]:
    r = HDR + day
    changed = []
    pairs = [
        ("B", vals.get("gas_vol")),
        ("C", vals.get("gas_profit")),
        ("J", vals.get("cstore_total")),
        ("K", vals.get("tax1", 0.0)),
        ("L", vals.get("tax4", 0.0)),
        ("M", vals.get("scratch", 0.0)),
        ("N", vals.get("lotto", 0.0)),
        ("O", vals.get("card", 0.0)),
    ]
    for col, val in pairs:
        if val is None:
            continue
        if ws[f"{col}{r}"].value in (None, ""):
            ws[f"{col}{r}"] = val
            changed.append(f"{col}={val}")
    if update_purch or purch is not None:
        ws[f"F{r}"] = purch
        changed.append(f"F=purch:{purch}")
    extras = {
        "D": f'=IF(OR(B{r}="",B{r}=0),"",C{r}/B{r})',
        "E": f'=IF(J{r}="","",J{r}-N(K{r})-N(L{r})-N(M{r})-N(N{r})-N(O{r}))',
        "G": f'=IF(OR(E{r}="",F{r}=""),"",E{r}-F{r})',
        "H": f'=IF(OR(E{r}="",E{r}=0),"",G{r}/E{r})',
        "I": f'=IF(OR(C{r}="",G{r}=""),"",C{r}+G{r})',
    }
    for col, formula in extras.items():
        if ws[f"{col}{r}"].value in (None, ""):
            ws[f"{col}{r}"] = formula
            changed.append(f"{col}=formula")
    return changed


def fill_san_diego(
    wb, day: int, vals: dict, purch: float | None = None, *, update_purch: bool = False
) -> list[str]:
    ws = wb["September Calculations"]
    r = 2 + day
    changed = []
    tax_sum = sum(
        (vals.get(k) or 0.0) for k in ("tax1", "tax4", "scratch", "lotto", "card")
    )
    mapping = [
        ("B", vals.get("gas_vol")),
        ("C", vals.get("gas_profit")),
        ("J", vals.get("cstore_total")),
        ("K", vals.get("tax1", 0.0)),
        ("L", vals.get("tax4", 0.0)),
        ("M", vals.get("scratch", 0.0)),
        ("N", vals.get("lotto", 0.0)),
        ("O", vals.get("card", 0.0)),
    ]
    if vals.get("cstore_total") is not None:
        mapping.append(("E", round(vals["cstore_total"] - tax_sum, 2)))
    for col, val in mapping:
        if val is None:
            continue
        if ws[f"{col}{r}"].value in (None, ""):
            ws[f"{col}{r}"] = val
            changed.append(f"{col}{r}={val}")
    if update_purch or purch is not None:
        ws[f"F{r}"] = purch
        changed.append(f"F{r}=purch:{purch}")
    return changed


def day_already_filled(wb, sid: str, day: int) -> bool:
    if sid == "42048":
        ws = wb["September Calculations"]
        return ws[f"B{2 + day}"].value not in (None, "")
    sheet = "September 2026"
    if sheet not in wb.sheetnames:
        return True
    ws = wb[sheet]
    cols = header_map(ws)
    gas_col = cols.get("gas_vol", "B")
    return ws[f"{gas_col}{HDR + day}"].value not in (None, "")


def fill_store(
    wb,
    sid: str,
    day: int,
    vals: dict,
    purch: float | None = None,
    *,
    update_purch: bool = False,
) -> list[str]:
    if sid == "42048":
        return fill_san_diego(wb, day, vals, purch=purch, update_purch=update_purch)
    if sid == "42098":
        return fill_brookhurst(wb, day, vals, purch=purch, update_purch=update_purch)
    if sid == "42399":
        return fill_garden_grove(
            wb["September 2026"], day, vals, purch=purch, update_purch=update_purch
        )
    ws = wb["September 2026"]
    cols = header_map(ws)
    return write_values(ws, day, vals, cols, purch=purch, update_purch=update_purch)


def write_purch_only(wb, sid: str, day: int, purch: float | None) -> list[str]:
    """Always overwrite Net Daily Purchases (None → blank). Sales untouched."""
    if sid == "42048":
        ws = wb["September Calculations"]
        r = 2 + day
        ws[f"F{r}"] = purch
        return [f"F{r}=purch:{purch}"]
    if sid == "42098":
        src = wb["September 2026 Source"]
        ws = wb["September 2026"]
        sr = 2 + day
        mr = HDR + day
        src[f"F{sr}"] = purch
        ws[f"F{mr}"] = f"='September 2026 Source'!F{sr}"
        return [f"src!F{sr}=purch:{purch}"]
    ws = wb["September 2026"]
    cols = header_map(ws)
    col = cols.get("purch") or "F"
    r = HDR + day
    ws[f"{col}{r}"] = purch
    return [f"{col}{r}=purch:{purch}"]


def parse_days_arg(raw: str) -> list[date]:
    out = []
    for part in raw.split(","):
        part = part.strip()
        if not part:
            continue
        out.append(date.fromisoformat(part))
    return out


def find_dly(dly_dir: Path, sub: str, through: date) -> Path | None:
    """Prefer stamped dly for through-date; else newest *dly.pdf in subdir."""
    direct = dly_dir / sub / dly_stamp(through)
    if direct.exists():
        return direct
    folder = dly_dir / sub
    if not folder.is_dir():
        return None
    cands = sorted(folder.glob("*dly.pdf"), key=lambda p: p.stat().st_mtime, reverse=True)
    return cands[0] if cands else None


def main() -> None:
    ap = argparse.ArgumentParser(description=__doc__)
    ap.add_argument("--xlsx-dir", required=True, help="Dir of * Daily.xlsx")
    ap.add_argument("--pdf-dir", required=True, help="Dir with Daily Book PDF subdirs")
    ap.add_argument(
        "--dly-dir",
        help="Dir with *dly.pdf subdirs (required for purchases)",
    )
    ap.add_argument(
        "--days",
        required=True,
        help="Comma YYYY-MM-DD days to fill sales (blank cells only)",
    )
    ap.add_argument(
        "--out-dir",
        help="Write filled workbooks here (default: modify xlsx-dir in place)",
    )
    ap.add_argument("--only", help="Comma store ids to limit")
    ap.add_argument(
        "--refresh-rules",
        action="store_true",
        help="Rebuild Deduct rules from Excel Deduct sheets before apply",
    )
    ap.add_argument(
        "--purchases-all-dly-days",
        action="store_true",
        help="Also write F for every Sep day that has Deduct nets in DLY "
        "(not only --days)",
    )
    ap.add_argument(
        "--report",
        default="/tmp/s2k/exports/fill_daily_excel_from_softsp.json",
    )
    args = ap.parse_args()

    xlsx_dir = Path(args.xlsx_dir)
    pdf_dir = Path(args.pdf_dir)
    dly_dir = Path(args.dly_dir) if args.dly_dir else None
    out_dir = Path(args.out_dir) if args.out_dir else xlsx_dir
    out_dir.mkdir(parents=True, exist_ok=True)
    days = parse_days_arg(args.days)
    through = max(days)
    only = {s.strip() for s in args.only.split(",")} if args.only else None

    # Purchase rules from Deduct sheets
    if args.refresh_rules:
        rules = scan_excel_dir(xlsx_dir)
        snap = REPO / "scripts" / "data" / "store_purchase_rules.json"
        snap.write_text(json.dumps(rules_to_json(rules), indent=2) + "\n")
        print(f"refreshed Deduct rules → {snap} ({len(rules)} stores)")
    else:
        rules = load_snapshot()

    # Parse DLY → per-store invoice nets
    purch_by_store: dict[str, dict[str, float | None]] = {}
    dly_invoice_count = 0
    if dly_dir:
        # Cache parsed BD dly once
        bd_inv: list[dict] | None = None
        for sid, meta in STORES.items():
            if only and sid not in only:
                continue
            _name, _fname, _mode, _dsub, _tso, dly_sub = meta
            rule = rules.get(sid) or load_deduct_from_workbook(
                xlsx_dir / _fname, store_id=sid
            )
            if dly_sub == "BD_central":
                if bd_inv is None:
                    p = find_dly(dly_dir, "BD_central", through)
                    if not p:
                        print(f"WARN no BD dly in {dly_dir}")
                        bd_inv = []
                    else:
                        bd_inv = parse_dly_invoices(p)
                        print(f"parsed BD dly {p.name}: {len(bd_inv)} lines")
                invs = [i for i in bd_inv if i["store_id"] == sid or i["store_id"].startswith(sid)]
            else:
                p = find_dly(dly_dir, dly_sub, through)
                if not p:
                    print(f"WARN no dly for {sid} in {dly_dir}/{dly_sub}")
                    invs = []
                else:
                    invs = parse_dly_invoices(p, {sid})
                    print(f"parsed {sid} dly {p.name}: {len(invs)} lines")
            dly_invoice_count += len(invs)
            purch_by_store[sid] = apply_invoices(rule, invs)
            # Show 22/23 nets
            for d in days:
                net = purch_by_store[sid].get(d.isoformat())
                print(f"  purch {sid} {d}: {net}")

    report = []
    for sid, (name, fname, mode, sub, tso, dly_sub) in STORES.items():
        if only and sid not in only:
            continue
        src = xlsx_dir / fname
        if not src.exists():
            report.append({"id": sid, "error": f"missing {src}"})
            print(f"SKIP {sid} missing {src}")
            continue
        dest = out_dir / fname
        if dest.resolve() != src.resolve():
            shutil.copy2(src, dest)
        wb = openpyxl.load_workbook(dest)
        filled = []
        purch_updates = []
        store_purch = purch_by_store.get(sid) or {}
        # When DLY was parsed for this store, always overwrite F (incl. blank)
        do_purch = sid in purch_by_store

        for d in days:
            if d.month != 9 or d.year != 2026:
                report.append(
                    {
                        "id": sid,
                        "date": d.isoformat(),
                        "error": "only September 2026 fill implemented",
                    }
                )
                continue
            day_n = d.day
            purch = store_purch.get(d.isoformat())
            # Missing key or Deduct-empty day → None → blank F when do_purch

            already = day_already_filled(wb, sid, day_n)
            pdf = pdf_dir / sub / pdf_stamp(d)
            vals: dict = {}
            if not already:
                if not pdf.exists():
                    print(f"  {sid} {d}: no Daily PDF {pdf}")
                    report.append({"id": sid, "date": d.isoformat(), "error": "no_pdf"})
                    if do_purch:
                        ch = write_purch_only(wb, sid, day_n, purch)
                        purch_updates.append({"date": d.isoformat(), "changed": ch})
                        print(f"  {sid} {d}: purch-only {ch}")
                    continue
                if mode == "single":
                    vals = parse_single(pdf)
                else:
                    vals = parse_bd_store(pdf, tso or sid) or {}
                    if not vals:
                        print(f"  {sid} {d}: TSO not in PDF")
                        report.append(
                            {"id": sid, "date": d.isoformat(), "error": "tso_miss"}
                        )
                        if do_purch:
                            ch = write_purch_only(wb, sid, day_n, purch)
                            purch_updates.append({"date": d.isoformat(), "changed": ch})
                            print(f"  {sid} {d}: purch-only {ch}")
                        continue
                if vals.get("gas_vol") is None or vals.get("cstore_total") is None:
                    print(f"  {sid} {d}: incomplete {vals}")
                    report.append(
                        {
                            "id": sid,
                            "date": d.isoformat(),
                            "error": "parse_incomplete",
                            "vals": vals,
                        }
                    )
                    if do_purch:
                        ch = write_purch_only(wb, sid, day_n, purch)
                        purch_updates.append({"date": d.isoformat(), "changed": ch})
                        print(f"  {sid} {d}: purch-only {ch}")
                    continue
                ch = fill_store(
                    wb, sid, day_n, vals, purch=purch, update_purch=do_purch
                )
            else:
                # Sales already present — still refresh purchases from DLY/Deduct
                if do_purch:
                    ch = write_purch_only(wb, sid, day_n, purch)
                    print(f"  {sid} {d}: sales already filled; purch refresh")
                else:
                    ch = []
                    print(f"  {sid} {d}: sales already filled; no dly purch")
            print(f"  {sid} {name} {d}: {ch}")
            filled.append(
                {
                    "date": d.isoformat(),
                    "changed": ch,
                    "purch": purch,
                    "vals": {
                        k: vals.get(k)
                        for k in (
                            "gas_vol",
                            "gas_profit",
                            "cstore_total",
                            "tax1",
                            "tax4",
                            "scratch",
                            "lotto",
                            "card",
                        )
                    }
                    if vals
                    else None,
                }
            )

        if args.purchases_all_dly_days and do_purch:
            for iso, net in store_purch.items():
                if not iso.startswith("2026-09"):
                    continue
                if any(f["date"] == iso for f in filled):
                    continue
                d = date.fromisoformat(iso)
                if not day_already_filled(wb, sid, d.day) and net is None:
                    continue
                ch = write_purch_only(wb, sid, d.day, net)
                if ch:
                    purch_updates.append({"date": iso, "changed": ch, "purch": net})

        if filled or purch_updates:
            wb.save(dest)
            print(
                f"SAVED {dest} sales_days={[f['date'] for f in filled]} "
                f"extra_purch={len(purch_updates)}"
            )
        else:
            print(f"no changes {sid}")
        report.append(
            {
                "id": sid,
                "name": name,
                "file": fname,
                "filled": [f["date"] for f in filled],
                "details": filled,
                "purch_extra": purch_updates,
            }
        )

    Path(args.report).parent.mkdir(parents=True, exist_ok=True)
    Path(args.report).write_text(
        json.dumps(
            {
                "through": through.isoformat(),
                "dly_invoice_lines": dly_invoice_count,
                "stations": report,
            },
            indent=2,
            default=str,
        )
        + "\n"
    )
    print(f"wrote {args.report}")


if __name__ == "__main__":
    main()
