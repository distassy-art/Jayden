#!/usr/bin/env python3
"""Fill Financial Audit S2K columns from Daily Book Summary PDFs.

Receipts mapping
----------------
  SAFEDROP                     → S2K Safe Drop
  CREDIT+DEBIT+EBT+MOBILE
    +PREPAID GIFT − FEE        → S2K Daily Receipt  (Credit Debit)
  abs(FEE)                     → EFT Fee Amount
  CASH OVER/SHORT              → Over / Short

Standing rule: always leave 2 days behind for the next round.
  Fill every available day through (PT today − 2); leave today and
  yesterday for the next run (e.g. Sep 13 → fill through Sep 11).

Example
-------
  python3 scripts/fill_financial_audit_from_s2k.py \\
    --xlsx /tmp/s2k/exports/placentia_financial_audit.xlsx \\
    --pdf-dir /tmp/s2k/pdfs/secondary_fill_d12 \\
    --station 42004 --year 2026 --month 9 \\
    --out /tmp/s2k/exports/placentia_financial_audit_filled.xlsx \\
    --upload --od-path "Clients/42004 (Arco Placentia)/Placentia Financial Audit.xlsx"
"""

from __future__ import annotations

import argparse
import json
import re
import ssl
import urllib.error
import urllib.parse
import urllib.request
from datetime import date, datetime, timedelta
from pathlib import Path
from typing import Any
from zoneinfo import ZoneInfo

import pdfplumber
from openpyxl import load_workbook
from openpyxl.worksheet.table import Table
from openpyxl.worksheet.worksheet import Worksheet

PT = ZoneInfo("America/Los_Angeles")
MONTH_NAMES = [
    "",
    "January",
    "February",
    "March",
    "April",
    "May",
    "June",
    "July",
    "August",
    "September",
    "October",
    "November",
    "December",
]

HOST = "https://smartsolutionsai26-my.sharepoint.com"
BASE = f"{HOST}/personal/minamorcos_smartsolutionsai26_onmicrosoft_com"
DOCS_ROOT = "/personal/minamorcos_smartsolutionsai26_onmicrosoft_com/Documents"
OD_COOKIES = Path("/tmp/od_cookies.json")
SSL_CTX = ssl.create_default_context()

LINE_RE = re.compile(
    r"^([A-Za-z][A-Za-z0-9 /&'.-]*?)\s+(\(\$?[\d,]+\.\d{2}\)|\$?[\d,]+\.\d{2})\s*$"
)


def parse_money(raw: str) -> float | None:
    t = raw.strip().replace(",", "").replace("$", "")
    if t.startswith("(") and t.endswith(")"):
        t = "-" + t[1:-1]
    if t in ("", "-", "--"):
        return None
    try:
        return float(t)
    except ValueError:
        return None


def parse_receipts(text: str) -> dict[str, float]:
    m = re.search(r"\nReceipts?\s*\n", text, re.I)
    block = text[m.start() :] if m else text
    out: dict[str, float] = {}
    for line in block.splitlines():
        line = line.strip()
        if not line:
            continue
        up = line.upper()
        if up.startswith(("RECEIPT TYPE", "TSO #", "LAST UPDATED")):
            continue
        mm = LINE_RE.match(line)
        if not mm:
            continue
        name = mm.group(1).strip().upper()
        amt = parse_money(mm.group(2))
        if amt is None or not name:
            continue
        out[name] = amt
    return out


def is_prepaid_gift(name: str) -> bool:
    compact = name.replace(" ", "")
    if compact in {"PREPAIDGIFT", "PREPAYGIFT", "PREPAIDGIFTCARD"}:
        return True
    return ("PREPAID" in name or "PREPAY" in name) and "GIFT" in name


def financial_fields(receipts: dict[str, float]) -> dict[str, float | None]:
    """Credit Debit = CREDIT+DEBIT+EBT+MOBILE+PREPAID GIFT − FEE."""
    safe = None
    for k, v in receipts.items():
        if k.replace(" ", "") in {"SAFEDROP", "SAFEDROP"}:
            safe = v
            break

    fee = None
    for k, v in receipts.items():
        if k == "FEE" or k.endswith(" FEE"):
            fee = abs(v)
            break

    over = None
    for k, v in receipts.items():
        if "OVER" in k and "SHORT" in k:
            over = v
            break

    daily = 0.0
    found = False
    for k, v in receipts.items():
        if (
            k in {"CREDIT", "DEBIT", "MOBILE"}
            or k.startswith("EBT")
            or is_prepaid_gift(k)
        ):
            daily += v
            found = True
    if found and fee is not None:
        daily -= fee

    return {
        "safe_drop": safe,
        "daily_receipt": round(daily, 2) if found else None,
        "eft_fee": fee,
        "over_short": over,
    }


def parse_daily_pdf(path: Path) -> dict[str, Any]:
    with pdfplumber.open(path) as pdf:
        text = "\n".join((p.extract_text() or "") for p in pdf.pages)
    receipts = parse_receipts(text)
    return {"path": str(path), "receipts": receipts, **financial_fields(receipts)}


def month_sheet_name(year: int, month: int) -> str:
    return f"{MONTH_NAMES[month]} {year}"


def default_through() -> date:
    """Last day to fill: PT today − 2 (leave 2 days for the next round)."""
    return datetime.now(PT).date() - timedelta(days=2)


def norm_col(name: str) -> str:
    return re.sub(r"[^a-z0-9]+", "", (name or "").lower())


def col_index(table: Table, *candidates: str) -> int | None:
    wanted = {norm_col(c) for c in candidates}
    for i, c in enumerate(table.tableColumns, start=1):
        if norm_col(c.name) in wanted:
            return i
    return None


def find_tables(ws: Worksheet) -> tuple[Table | None, Table | None]:
    cashier = over = None
    for t in ws.tables.values():
        cols = {norm_col(c.name) for c in t.tableColumns}
        if "s2ksafedrop" in cols:
            cashier = t
        if "overshort" in cols:
            over = t
    return cashier, over


def parse_ref(ref: str) -> tuple[int, int]:
    m = re.match(r"[A-Z]+(\d+):[A-Z]+(\d+)", ref)
    if not m:
        raise ValueError(f"Bad table ref {ref}")
    return int(m.group(1)), int(m.group(2))


def as_date(v: Any) -> date | None:
    if isinstance(v, datetime):
        return v.date()
    if isinstance(v, date):
        return v
    return None


def fix_broken_split_refs(ws: Worksheet) -> int:
    n = 0
    for row in ws.iter_rows(min_row=1, max_row=min(ws.max_row or 1, 120), max_col=12):
        for cell in row:
            v = cell.value
            if not isinstance(v, str):
                continue
            nv = v.replace("$B$224", "$B$7")
            if cell.coordinate in {"H7", "H8", "E7", "F7", "G7", "E8", "F8", "G8"}:
                nv = (
                    nv.replace("E224", "E7")
                    .replace("F224", "F7")
                    .replace("G224", "G7")
                    .replace("E225", "E8")
                    .replace("F225", "F8")
                    .replace("G225", "G8")
                )
            if nv != v:
                cell.value = nv
                n += 1
    for t in ws.tables.values():
        for c in t.tableColumns:
            formula = getattr(c, "calculatedColumnFormula", None)
            if not formula:
                continue
            text = getattr(formula, "text", None)
            if isinstance(text, str) and "$B$224" in text:
                formula.text = text.replace("$B$224", "$B$7")
                n += 1
    return n


def fill_month(
    ws: Worksheet,
    year: int,
    month: int,
    by_day: dict[int, dict[str, Any]],
    overwrite: bool = True,
) -> dict[str, Any]:
    cashier, over = find_tables(ws)
    if cashier is None:
        raise SystemExit(f"No cashier table on {ws.title}")

    header_row, last_row = parse_ref(cashier.ref)
    has_totals = any(
        bool(c.totalsRowFunction or c.totalsRowLabel or c.totalsRowFormula)
        for c in cashier.tableColumns
    )
    data_last = last_row - (1 if has_totals else 0)

    c_date = col_index(cashier, "Date") or 1
    c_safe = col_index(cashier, "S2K Safe Drop")
    c_daily = col_index(cashier, "S2K Daily Receipt")
    c_fee = col_index(cashier, "EFT Fee Amount")
    if not all([c_safe, c_daily, c_fee]):
        names = [c.name for c in cashier.tableColumns]
        raise SystemExit(f"Missing S2K cols on {ws.title}; have {names}")

    date_rows: dict[date, int] = {}
    empty_rows: list[int] = []
    for r in range(header_row + 1, data_last + 1):
        d = as_date(ws.cell(r, c_date).value)
        if d:
            date_rows[d] = r
        elif ws.cell(r, c_date).value in (None, ""):
            empty_rows.append(r)

    filled: list[dict[str, Any]] = []
    skipped: list[dict[str, Any]] = []
    for day, fields in sorted(by_day.items()):
        d = date(year, month, day)
        row = date_rows.get(d)
        if row is None:
            if not empty_rows:
                skipped.append({"day": day, "reason": "no empty row"})
                continue
            row = empty_rows.pop(0)
            ws.cell(row, c_date).value = datetime(year, month, day)
            date_rows[d] = row

        if not overwrite and ws.cell(row, c_safe).value not in (None, ""):
            skipped.append({"day": day, "reason": "already filled"})
            continue

        if fields.get("safe_drop") is not None:
            ws.cell(row, c_safe).value = fields["safe_drop"]
        if fields.get("daily_receipt") is not None:
            ws.cell(row, c_daily).value = round(float(fields["daily_receipt"]), 2)
        if fields.get("eft_fee") is not None:
            ws.cell(row, c_fee).value = fields["eft_fee"]
        filled.append(
            {
                "day": day,
                "row": row,
                "safe_drop": fields.get("safe_drop"),
                "daily_receipt": fields.get("daily_receipt"),
                "eft_fee": fields.get("eft_fee"),
                "over_short": fields.get("over_short"),
            }
        )

    over_filled: list[dict[str, Any]] = []
    if over is not None:
        o_header, o_last = parse_ref(over.ref)
        o_has_totals = any(
            bool(c.totalsRowFunction or c.totalsRowLabel or c.totalsRowFormula)
            for c in over.tableColumns
        )
        o_data_last = o_last - (1 if o_has_totals else 0)
        header_val = ws.cell(o_header, 1).value
        first_data = o_header + (1 if isinstance(header_val, str) else 0)

        o_date_rows: dict[date, int] = {}
        o_empty: list[int] = []
        for r in range(first_data, o_data_last + 1):
            d = as_date(ws.cell(r, 1).value)
            if d:
                o_date_rows[d] = r
            elif ws.cell(r, 1).value in (None, ""):
                o_empty.append(r)

        for day, fields in sorted(by_day.items()):
            if fields.get("over_short") is None:
                continue
            d = date(year, month, day)
            row = o_date_rows.get(d)
            if row is None:
                if not o_empty:
                    continue
                row = o_empty.pop(0)
                ws.cell(row, 1).value = datetime(year, month, day)
                o_date_rows[d] = row
            if overwrite or ws.cell(row, 2).value in (None, ""):
                ws.cell(row, 2).value = fields["over_short"]
                over_filled.append(
                    {"day": day, "row": row, "over_short": fields["over_short"]}
                )

    refs_fixed = fix_broken_split_refs(ws)
    return {
        "sheet": ws.title,
        "cashier_filled": filled,
        "over_filled": over_filled,
        "skipped": skipped,
        "refs_fixed": refs_fixed,
    }


def cookie_header() -> str:
    cookies = json.loads(OD_COOKIES.read_text())
    return "; ".join(
        f"{c['name']}={c['value']}"
        for c in cookies
        if "sharepoint.com" in c.get("domain", "")
    )


def od_req(
    url: str,
    method: str = "GET",
    data: bytes | None = None,
    headers: dict | None = None,
) -> tuple[int, bytes]:
    h = {
        "Cookie": cookie_header(),
        "Accept": "application/json;odata=verbose",
        "User-Agent": "Mozilla/5.0",
    }
    if headers:
        h.update(headers)
    req = urllib.request.Request(url, data=data, headers=h, method=method)
    try:
        with urllib.request.urlopen(req, context=SSL_CTX, timeout=180) as resp:
            return resp.status, resp.read()
    except urllib.error.HTTPError as e:
        return e.code, e.read()


def digest() -> str:
    return json.loads(od_req(BASE + "/_api/contextinfo", method="POST", data=b"")[1])[
        "d"
    ]["GetContextWebInformation"]["FormDigestValue"]


def upload_od(rel_under_docs: str, content: bytes) -> bool:
    server_rel = f"{DOCS_ROOT}/{rel_under_docs.lstrip('/')}"
    folder, name = server_rel.rsplit("/", 1)
    enc = urllib.parse.quote(server_rel, safe="/")
    dig = digest()
    for action in [
        "CheckOut",
        "UndoCheckOut",
        "CheckIn(comment='financial-audit-s2k-fill',checkInType=1)",
        "UndoCheckOut",
    ]:
        od_req(
            BASE + f"/_api/web/GetFileByServerRelativeUrl('{enc}')/{action}",
            method="POST",
            data=b"",
            headers={"X-RequestDigest": dig, "IF-MATCH": "*"},
        )
    add_url = (
        BASE
        + f"/_api/web/GetFolderByServerRelativeUrl('{urllib.parse.quote(folder, safe='/')}')"
        + f"/Files/add(url='{urllib.parse.quote(name)}',overwrite=true)"
    )
    for i in range(6):
        dig = digest()
        code, body = od_req(
            add_url,
            method="POST",
            data=content,
            headers={
                "X-RequestDigest": dig,
                "Content-Type": "application/octet-stream",
                "Prefer": "bypass-shared-lock",
            },
        )
        if code in (200, 201):
            return True
        print(f"  upload retry {i + 1}: {code} {body[:160]!r}")
    return False


def discover_pdfs(
    pdf_dir: Path, station: str, year: int, month: int, through: date
) -> dict[int, Path]:
    out: dict[int, Path] = {}
    patterns = [
        re.compile(rf"^{re.escape(station)}_(\d{{2}})(\d{{2}}){year}\.pdf$", re.I),
        re.compile(rf"^(\d{{2}})(\d{{2}}){year}\.pdf$", re.I),
    ]
    for p in sorted(pdf_dir.glob("*.pdf")):
        if re.search(r"dly|dpt", p.name, re.I):
            continue
        for pat in patterns:
            m = pat.match(p.name)
            if not m:
                continue
            mm, dd = int(m.group(1)), int(m.group(2))
            if mm != month:
                continue
            d = date(year, month, dd)
            if d > through:
                continue
            out[dd] = p
            break
    return out


def main() -> None:
    ap = argparse.ArgumentParser(
        description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter
    )
    ap.add_argument("--xlsx", required=True, type=Path)
    ap.add_argument("--pdf-dir", required=True, type=Path)
    ap.add_argument("--station", required=True)
    ap.add_argument("--year", type=int, required=True)
    ap.add_argument("--month", type=int, required=True)
    ap.add_argument(
        "--through",
        type=date.fromisoformat,
        default=None,
        help="Last day to fill (YYYY-MM-DD). Default: PT today−2 (leave 2 days behind).",
    )
    ap.add_argument("--out", type=Path, help="Output xlsx (default: overwrite --xlsx)")
    ap.add_argument(
        "--no-overwrite",
        action="store_true",
        help="Skip days that already have S2K Safe Drop",
    )
    ap.add_argument("--upload", action="store_true")
    ap.add_argument(
        "--od-path",
        default="Clients/42004 (Arco Placentia)/Placentia Financial Audit.xlsx",
    )
    ap.add_argument("--report", type=Path, help="Write JSON report")
    args = ap.parse_args()

    through = args.through or default_through()
    print(
        f"Fill through={through.isoformat()} "
        f"(leave 2 days behind; PT today−2 unless --through set)"
    )

    pdfs = discover_pdfs(
        args.pdf_dir, args.station, args.year, args.month, through=through
    )
    if not pdfs:
        raise SystemExit(
            f"No daily PDFs for {args.station} {args.year}-{args.month:02d} "
            f"through {through} in {args.pdf_dir}"
        )

    by_day: dict[int, dict[str, Any]] = {}
    parse_errors: list[dict[str, Any]] = []
    for day, path in sorted(pdfs.items()):
        try:
            parsed = parse_daily_pdf(path)
            if parsed.get("safe_drop") is None and parsed.get("daily_receipt") is None:
                parse_errors.append(
                    {"day": day, "path": str(path), "reason": "no receipts parsed"}
                )
                continue
            by_day[day] = parsed
        except Exception as e:  # noqa: BLE001
            parse_errors.append({"day": day, "path": str(path), "error": str(e)})

    wb = load_workbook(args.xlsx)
    sheet = month_sheet_name(args.year, args.month)
    if sheet not in wb.sheetnames:
        raise SystemExit(f"Missing sheet {sheet}; have {wb.sheetnames}")

    result = fill_month(
        wb[sheet],
        args.year,
        args.month,
        by_day,
        overwrite=not args.no_overwrite,
    )
    for sn in wb.sheetnames:
        if sn == sheet:
            continue
        if re.search(r"20\d{2}", sn):
            result.setdefault("other_refs_fixed", {})[sn] = fix_broken_split_refs(
                wb[sn]
            )

    out = args.out or args.xlsx
    wb.save(out)

    report = {
        "station": args.station,
        "year": args.year,
        "month": args.month,
        "through": through.isoformat(),
        "pdfs_found": {str(k): str(v) for k, v in pdfs.items()},
        "parsed_days": sorted(by_day),
        "parse_errors": parse_errors,
        "fill": result,
        "out": str(out),
        "upload": None,
    }

    if args.upload:
        ok = upload_od(args.od_path, out.read_bytes())
        report["upload"] = {"ok": ok, "path": args.od_path}
        print("UPLOADED" if ok else "UPLOAD FAILED", args.od_path)

    if args.report:
        args.report.parent.mkdir(parents=True, exist_ok=True)
        args.report.write_text(json.dumps(report, indent=2, default=str))

    print(
        f"Filled {sheet} through {through}: "
        f"{len(result['cashier_filled'])} cashier days, "
        f"{len(result['over_filled'])} over/short, refs_fixed={result['refs_fixed']}"
    )
    for row in result["cashier_filled"]:
        print(
            f"  {args.month:02d}/{row['day']:02d}: safe={row['safe_drop']} "
            f"daily={row['daily_receipt']} fee={row['eft_fee']} os={row['over_short']}"
        )


if __name__ == "__main__":
    main()
