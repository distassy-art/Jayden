#!/usr/bin/env python3
"""Split combined client Audit.xlsx into Vendor + Financial workbooks.

Vendor Audit: scanned invoices + active S2K + electronic/excluded S2K
(+ Documents helper sheets when present).

Financial Audit: cashier safe-drop + cash over/short + EOM safe-drop.

Original combined ``* Audit.xlsx`` is left in place as archive.

Usage:
  python3 scripts/split_audit_workbook.py --local /path/to/Placentia\\ Audit.xlsx
  python3 scripts/split_audit_workbook.py --all-clients
  python3 scripts/split_audit_workbook.py --station 42179 --station 42004
"""
from __future__ import annotations

import argparse
import json
import re
import ssl
import urllib.error
import urllib.parse
import urllib.request
from copy import copy as shallow_copy
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.worksheet.table import Table, TableColumn, TableStyleInfo

HOST = "https://smartsolutionsai26-my.sharepoint.com"
BASE = f"{HOST}/personal/minamorcos_smartsolutionsai26_onmicrosoft_com"
DOCS = "/personal/minamorcos_smartsolutionsai26_onmicrosoft_com/Documents"
COOKIES_PATH = Path("/tmp/od_cookies.json")
CLIENTS_JSON = Path(__file__).resolve().parent / "invoice_scan_clients.json"
SSL_CTX = ssl.create_default_context()
CASHIER_TITLE = "CASHIER SAFE DROP AUDIT"


def cookie_header() -> str:
    cookies = json.loads(COOKIES_PATH.read_text())
    return "; ".join(
        f"{c['name']}={c['value']}"
        for c in cookies
        if "sharepoint.com" in c.get("domain", "")
    )


def http(
    url: str,
    method: str = "GET",
    data: bytes | None = None,
    headers: dict | None = None,
    timeout: int = 180,
):
    h = {
        "Cookie": cookie_header(),
        "Accept": "application/json;odata=verbose",
        "User-Agent": "Mozilla/5.0",
    }
    if headers:
        h.update(headers)
    req = urllib.request.Request(url, data=data, headers=h, method=method)
    try:
        with urllib.request.urlopen(req, context=SSL_CTX, timeout=timeout) as resp:
            return resp.status, resp.read()
    except urllib.error.HTTPError as e:
        return e.code, e.read()


def digest() -> str:
    code, body = http(BASE + "/_api/contextinfo", method="POST", data=b"")
    if code != 200:
        raise SystemExit(f"digest fail {code} {body[:200]}")
    return json.loads(body)["d"]["GetContextWebInformation"]["FormDigestValue"]


def list_files(folder_rel: str) -> list[dict]:
    url = (
        BASE
        + f"/_api/web/GetFolderByServerRelativeUrl('{urllib.parse.quote(folder_rel)}')"
        + "/Files?$select=Name,Length,TimeLastModified,ServerRelativeUrl&$top=200"
    )
    code, body = http(url)
    if code != 200:
        print(f"  list fail {folder_rel} {code} {body[:120]}")
        return []
    return json.loads(body)["d"]["results"]


def download(server_rel: str) -> bytes:
    url = (
        BASE
        + f"/_api/web/GetFileByServerRelativeUrl('{urllib.parse.quote(server_rel)}')/$value"
    )
    code, body = http(url, timeout=300)
    if code != 200:
        raise RuntimeError(f"download {code} {server_rel} {body[:160]}")
    return body


def upload(folder_rel: str, name: str, data: bytes) -> bool:
    url = (
        BASE
        + f"/_api/web/GetFolderByServerRelativeUrl('{urllib.parse.quote(folder_rel)}')"
        + f"/Files/add(url='{urllib.parse.quote(name)}',overwrite=true)"
    )
    code, body = http(
        url,
        method="POST",
        data=data,
        headers={
            "X-RequestDigest": digest(),
            "Content-Type": "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        },
        timeout=300,
    )
    if code in (200, 201):
        return True
    print(f"  upload fail {name} {code} {body[:200]}")
    return False


def safe_set(ws, coord: str, value) -> None:
    """Set a cell value even if it sits in a merged range."""
    cell = ws[coord]
    # If this is a MergedCell, unmerge the range that contains it first.
    from openpyxl.cell.cell import MergedCell

    if isinstance(cell, MergedCell):
        for m in list(ws.merged_cells.ranges):
            if coord in m:
                ws.unmerge_cells(str(m))
                break
        cell = ws[coord]
    cell.value = value


def find_cashier_start(ws) -> int | None:
    for r in range(1, (ws.max_row or 0) + 1):
        for c in range(1, 6):
            v = ws.cell(r, c).value
            if isinstance(v, str) and v.strip().upper() == CASHIER_TITLE:
                return r
    return None


def table_kind(name: str) -> str:
    n = name.lower()
    if "cashier" in n or "cashover" in n or "overshort" in n:
        return "fin"
    if "scanned" in n or "excluded" in n or "s2k" in n or "document" in n:
        return "vendor"
    return "other"


def copy_cell(src, dst) -> None:
    dst.value = src.value
    if src.has_style:
        dst.font = shallow_copy(src.font)
        dst.border = shallow_copy(src.border)
        dst.fill = shallow_copy(src.fill)
        dst.number_format = src.number_format
        dst.protection = shallow_copy(src.protection)
        dst.alignment = shallow_copy(src.alignment)


def copy_col_dims(sw, dw) -> None:
    for letter, dim in sw.column_dimensions.items():
        dw.column_dimensions[letter].width = dim.width
        dw.column_dimensions[letter].hidden = dim.hidden


def clone_table(tbl, new_ref: str) -> Table:
    nt = Table(displayName=tbl.displayName, ref=new_ref)
    nt.name = tbl.name
    cols = []
    for c in tbl.tableColumns:
        tc = TableColumn(id=c.id, name=c.name)
        if c.totalsRowLabel is not None:
            tc.totalsRowLabel = c.totalsRowLabel
        if c.totalsRowFunction is not None:
            tc.totalsRowFunction = c.totalsRowFunction
        if c.totalsRowFormula is not None:
            tc.totalsRowFormula = shallow_copy(c.totalsRowFormula)
        if c.calculatedColumnFormula is not None:
            tc.calculatedColumnFormula = shallow_copy(c.calculatedColumnFormula)
        cols.append(tc)
    nt.tableColumns = cols
    if tbl.totalsRowCount:
        nt.totalsRowCount = tbl.totalsRowCount
    if tbl.totalsRowShown is not None:
        nt.totalsRowShown = tbl.totalsRowShown
    if tbl.headerRowCount is not None:
        nt.headerRowCount = tbl.headerRowCount
    if tbl.tableStyleInfo is not None:
        s = tbl.tableStyleInfo
        nt.tableStyleInfo = TableStyleInfo(
            name=s.name,
            showFirstColumn=s.showFirstColumn,
            showLastColumn=s.showLastColumn,
            showRowStripes=s.showRowStripes,
            showColumnStripes=s.showColumnStripes,
        )
    return nt


def scrub_cashier_formulas(ws) -> None:
    for row in ws.iter_rows(
        min_row=1, max_row=ws.max_row or 1, max_col=ws.max_column or 1
    ):
        for cell in row:
            v = cell.value
            if isinstance(v, str) and v.startswith("=") and re.search(
                r"Cashier|CashOverShort", v, re.I
            ):
                cell.value = None
            if isinstance(v, str) and v.strip().upper() in (
                "SAFE DROP VARIANCE",
                "SAFE DROP VAR",
            ):
                cell.value = None


def split_names(combined_name: str) -> tuple[str, str]:
    """'Placentia Audit.xlsx' -> Vendor / Financial names."""
    lower = combined_name.lower()
    idx = lower.rfind(" audit.xlsx")
    if idx >= 0:
        base = combined_name[:idx]
        return f"{base} Vendor Audit.xlsx", f"{base} Financial Audit.xlsx"
    stem = Path(combined_name).stem
    return f"{stem} Vendor Audit.xlsx", f"{stem} Financial Audit.xlsx"


def is_documents_sheet(name: str) -> bool:
    return "document" in name.lower()


def build_vendor(src_path: Path, out_path: Path, store_label: str) -> dict:
    wb = load_workbook(src_path)
    if "Cover" in wb.sheetnames:
        cover = wb["Cover"]
        safe_set(cover, "A1", "2026 Store Vendor Audit Workbook")
        safe_set(
            cover,
            "A2",
            f"Monthly scanned-invoice vs S2K vendor reconciliation — {store_label}",
        )
        safe_set(
            cover,
            "A12",
            "VENDOR AUDIT — Scanned invoices + Active S2K + Electronic/Excluded S2K. "
            "Cashier / cash O-S / EOM safe-drop live in the Financial Audit workbook. "
            "Original combined Audit.xlsx is kept as archive.",
        )

    report: dict = {"sheets": {}}
    for sn in list(wb.sheetnames):
        if sn == "Cover":
            continue
        ws = wb[sn]
        if is_documents_sheet(sn):
            report["sheets"][sn] = "kept (documents)"
            continue
        start = find_cashier_start(ws)
        for name in list(ws.tables):
            if table_kind(name) == "fin":
                del ws.tables[name]
        if start:
            ws.delete_rows(start, (ws.max_row or start) - start + 1)
        scrub_cashier_formulas(ws)
        try:
            a2 = ws["A2"].value
        except Exception:
            a2 = None
        if isinstance(a2, str) and "cashier" in a2.lower():
            safe_set(ws, "A2", "Vendor invoice comparison: scanned invoices vs S2K")
        report["sheets"][sn] = {
            "cashier_cut": start,
            "tables": list(ws.tables),
            "rows": ws.max_row,
        }

    wb.save(out_path)
    return report


def build_financial(src_path: Path, out_path: Path, store_label: str) -> dict:
    src = load_workbook(src_path)
    out = Workbook()
    cover = out.active
    cover.title = "Cover"
    if "Cover" in src.sheetnames:
        sc = src["Cover"]
        for r in range(1, (sc.max_row or 1) + 1):
            for c in range(1, (sc.max_column or 1) + 1):
                copy_cell(sc.cell(r, c), cover.cell(r, c))
        copy_col_dims(sc, cover)
    # Drop copied Cover merges so title cells are writable.
    for m in list(cover.merged_cells.ranges):
        cover.unmerge_cells(str(m))
    safe_set(cover, "A1", "2026 Store Financial Audit Workbook")
    safe_set(
        cover,
        "A2",
        f"Cashier safe-drop, electronic tender, and cash over/short — {store_label}",
    )
    safe_set(
        cover,
        "A12",
        "FINANCIAL AUDIT — Cashier safe drop + Cash over/short + EOM safe-drop. "
        "Scanned + S2K vendor sections live in the Vendor Audit workbook. "
        "Original combined Audit.xlsx is kept as archive.",
    )

    report: dict = {"sheets": {}}
    for sn in src.sheetnames:
        if sn == "Cover" or is_documents_sheet(sn):
            continue
        sw = src[sn]
        start = find_cashier_start(sw)
        if not start:
            report["sheets"][sn] = "skipped (no cashier section)"
            continue
        dw = out.create_sheet(sn)
        dw["A1"].value = f"{sn} Financial Audit"
        dw["A2"].value = "Cashier safe-drop / electronic / cash over-short"
        dw["A3"].value = f"Split from combined Audit — {store_label}"
        copy_col_dims(sw, dw)
        dest_offset = 5 - start
        max_r = sw.max_row or start
        max_c = min(sw.max_column or 9, 12)
        for r in range(start, max_r + 1):
            dr = r + dest_offset
            if r in sw.row_dimensions and sw.row_dimensions[r].height:
                dw.row_dimensions[dr].height = sw.row_dimensions[r].height
            for c in range(1, max_c + 1):
                copy_cell(sw.cell(r, c), dw.cell(dr, c))
        for m in sw.merged_cells.ranges:
            if m.max_row < start:
                continue
            min_r = max(m.min_row, start)
            dw.merge_cells(
                start_row=min_r + dest_offset,
                start_column=m.min_col,
                end_row=m.max_row + dest_offset,
                end_column=m.max_col,
            )
        for name in list(sw.tables):
            if table_kind(name) != "fin":
                continue
            tbl = sw.tables[name]
            m = re.match(r"([A-Z]+)(\d+):([A-Z]+)(\d+)", tbl.ref)
            if not m:
                continue
            a, r1, b, r2 = m.group(1), int(m.group(2)), m.group(3), int(m.group(4))
            new_ref = f"{a}{r1 + dest_offset}:{b}{r2 + dest_offset}"
            try:
                dw.add_table(clone_table(tbl, new_ref))
            except Exception as e:
                print(f"  table {name} skip: {e}")
        report["sheets"][sn] = {
            "cashier_start": start,
            "tables": list(dw.tables),
            "rows": dw.max_row,
        }

    out.save(out_path)
    return report


def find_combined_audit(folder_rel: str) -> str | None:
    files = list_files(folder_rel)
    names = [f["Name"] for f in files if f["Name"].lower().endswith(".xlsx")]
    audits = [n for n in names if "audit" in n.lower()]
    combined = [
        n
        for n in audits
        if "vendor audit" not in n.lower()
        and "financial audit" not in n.lower()
        and "prototype" not in n.lower()
        and "combined" not in n.lower()
    ]
    preferred = [n for n in combined if n.lower().endswith(" audit.xlsx")]
    if preferred:
        return sorted(preferred, key=len)[0]
    return combined[0] if combined else None


def process_client(dest: str, name: str, station: str | None, out_dir: Path) -> dict:
    folder_rel = f"{DOCS}/{dest}"
    combined = find_combined_audit(folder_rel)
    if not combined:
        return {"client": name, "station": station, "skipped": "no combined Audit.xlsx"}
    vendor_name, fin_name = split_names(combined)
    label = f"{station or ''} {name}".strip()
    print(f"\n=== {label} :: {combined} ===")
    raw = download(f"{folder_rel}/{combined}")
    src_path = out_dir / f"{station or name}_combined.xlsx"
    vendor_path = out_dir / vendor_name
    fin_path = out_dir / fin_name
    src_path.write_bytes(raw)
    vrep = build_vendor(src_path, vendor_path, label)
    frep = build_financial(src_path, fin_path, label)
    ok_v = upload(folder_rel, vendor_name, vendor_path.read_bytes())
    ok_f = upload(folder_rel, fin_name, fin_path.read_bytes())
    print(f"  upload vendor={ok_v} financial={ok_f}")
    for f in list_files(folder_rel):
        n = f["Name"]
        if "prototype" in n.lower() and "audit" in n.lower():
            print(f"  note: leftover prototype {n}")
    return {
        "client": name,
        "station": station,
        "combined": combined,
        "vendor": vendor_name,
        "financial": fin_name,
        "upload_vendor": ok_v,
        "upload_financial": ok_f,
        "vendor_report": vrep,
        "financial_report": frep,
    }


def main() -> int:
    p = argparse.ArgumentParser(description=__doc__)
    p.add_argument("--local", help="Split a local combined Audit.xlsx")
    p.add_argument("--all-clients", action="store_true")
    p.add_argument("--station", action="append", help="Limit to station id(s)")
    p.add_argument("--dry-run", action="store_true")
    args = p.parse_args()

    out_dir = Path("/tmp/s2k/exports/audit_split")
    out_dir.mkdir(parents=True, exist_ok=True)

    if args.local:
        src = Path(args.local)
        vname, fname = split_names(src.name)
        label = src.stem
        build_vendor(src, src.with_name(vname), label)
        build_financial(src, src.with_name(fname), label)
        print("wrote", src.with_name(vname))
        print("wrote", src.with_name(fname))
        return 0

    if not args.all_clients and not args.station:
        p.error("use --all-clients, --station, or --local")

    cfg = json.loads(CLIENTS_JSON.read_text())
    clients = cfg["clients"]
    if args.station:
        want = set(args.station)
        clients = [c for c in clients if str(c.get("station") or "") in want]

    results = []
    for c in clients:
        dest = c["dest"]
        name = c["name"]
        station = str(c.get("station") or "")
        if args.dry_run:
            combined = find_combined_audit(f"{DOCS}/{dest}")
            print(station, name, "->", combined or "NO AUDIT")
            continue
        try:
            results.append(process_client(dest, name, station, out_dir))
        except Exception as e:
            print(f"  FAIL {name}: {e}")
            results.append({"client": name, "station": station, "error": str(e)})

    if results:
        report_path = out_dir / "split_report.json"
        report_path.write_text(json.dumps(results, indent=2, default=str))
        print("\nReport", report_path)
        ok = sum(
            1 for r in results if r.get("upload_vendor") and r.get("upload_financial")
        )
        print(f"Uploaded both files for {ok}/{len(results)} clients")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
