#!/usr/bin/env python3
"""Build one vendor-spend Excel per client: monthly sheets + yearly summary."""
from __future__ import annotations

import re
from collections import defaultdict
from pathlib import Path

import pdfplumber
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

MONEY = re.compile(r"-?\$[\d,]+\.\d{2}")
SKIP = re.compile(
    r"(?i)^(station\s+total|sales\s+tax|page\s+total|grand\s+total|report\s+total|"
    r"grouped by|station$|inv\s*#|1\s+station|\d+\s+of\s+\d+)"
)
THIN = Border(
    left=Side(style="thin"),
    right=Side(style="thin"),
    top=Side(style="thin"),
    bottom=Side(style="thin"),
)
HEADER_FONT = Font(bold=True, color="FFFFFF")
HEADER_FILL = PatternFill("solid", fgColor="1F4E79")


def parse_money(s) -> float:
    if s is None or s == "":
        return 0.0
    s = str(s).strip().replace("$", "").replace(",", "")
    try:
        return float(s)
    except ValueError:
        return 0.0


def parse_pdf_tables(path: Path):
    """Vendor spend uses Inv Total (col 3). Invoice lines start with TSO."""
    vendors = []
    cur = None
    grand_inv = None
    with pdfplumber.open(path) as pdf:
        for page in pdf.pages:
            for table in page.extract_tables() or []:
                for row in table:
                    if not row:
                        continue
                    joined = " ".join(str(c or "") for c in row)
                    if "Grand Total" in joined:
                        vals = []
                        for c in row:
                            if not c:
                                continue
                            cs = str(c).strip()
                            if MONEY.search(cs) or re.match(r"^-?[\d,]+\.\d{2}$", cs):
                                vals.append(parse_money(cs))
                        if vals:
                            grand_inv = vals[0]
                        continue
                    if not row[0] or not str(row[0]).strip():
                        continue
                    name = str(row[0]).strip()
                    if SKIP.search(name):
                        continue
                    if name.upper().startswith("TSO"):
                        if cur is not None:
                            cur["orders"] += 1
                        continue
                    inv = parse_money(row[3] if len(row) > 3 else None)
                    is_vendor = len(row) >= 4 and (row[1] in (None, "")) and re.search(r"[A-Za-z]", name)
                    if is_vendor:
                        cur = {
                            "vendor": re.sub(r"\s+", " ", name),
                            "total": inv,
                            "orders": 0,
                        }
                        vendors.append(cur)
    return vendors, grand_inv


def month_label(mm: str) -> str:
    names = ["", "Jan", "Feb", "Mar", "Apr", "May", "Jun", "Jul", "Aug", "Sep", "Oct", "Nov", "Dec"]
    return f"{names[int(mm)]} 2026"


def write_vendor_sheet(ws, title, subtitle, rows, period_total, pct_label, extra_col=None):
    ws["A1"] = title
    ws["A1"].font = Font(bold=True, size=14)
    last_col = "F" if extra_col else "E"
    ws.merge_cells(f"A1:{last_col}1")
    ws["A2"] = subtitle
    ws["A2"].font = Font(bold=True)
    headers = ["Vendor", "Orders", "Total Spend", "Avg per Order", pct_label]
    if extra_col:
        headers.append(extra_col)
    for col, h in enumerate(headers, 1):
        cell = ws.cell(4, col, h)
        cell.font = HEADER_FONT
        cell.fill = HEADER_FILL
        cell.alignment = Alignment(horizontal="center")
        cell.border = THIN
    for i, r in enumerate(rows, 5):
        ws.cell(i, 1, r["vendor"]).border = THIN
        c2 = ws.cell(i, 2, r["orders"])
        c2.border = THIN
        c2.alignment = Alignment(horizontal="center")
        c3 = ws.cell(i, 3, round(r["total"], 2))
        c3.border = THIN
        c3.number_format = "#,##0.00"
        avg = (r["total"] / r["orders"]) if r["orders"] else None
        c4 = ws.cell(i, 4, round(avg, 2) if avg is not None else "")
        c4.border = THIN
        if avg is not None:
            c4.number_format = "#,##0.00"
        pct = (r["total"] / period_total) if period_total else 0
        c5 = ws.cell(i, 5, pct)
        c5.border = THIN
        c5.number_format = "0.00%"
        if extra_col:
            ws.cell(i, 6, r.get("months", "")).border = THIN
    total_row = 4 + len(rows) + 1
    ws.cell(total_row, 1, "TOTAL").font = Font(bold=True)
    ws.cell(total_row, 2, sum(r["orders"] for r in rows)).font = Font(bold=True)
    c = ws.cell(total_row, 3, round(period_total, 2))
    c.font = Font(bold=True)
    c.number_format = "#,##0.00"
    c5 = ws.cell(total_row, 5, 1 if period_total else 0)
    c5.number_format = "0.00%"
    c5.font = Font(bold=True)
    widths = [48, 10, 14, 14, 16] + ([14] if extra_col else [])
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = w


def build_client_workbook(client_name: str, monthly: dict, out_path: Path):
    wb = Workbook()
    months = [m for m in sorted(monthly.keys()) if monthly[m]]
    vendor_agg = defaultdict(lambda: {"orders": 0, "total": 0.0, "months": set()})
    year_total = 0.0
    for mm in months:
        for r in monthly[mm]:
            a = vendor_agg[r["vendor"]]
            a["orders"] += r["orders"]
            a["total"] += r["total"]
            a["months"].add(mm)
            year_total += r["total"]
    year_rows = sorted(
        [
            {
                "vendor": v,
                "orders": a["orders"],
                "total": a["total"],
                "months": len(a["months"]),
            }
            for v, a in vendor_agg.items()
        ],
        key=lambda x: -x["total"],
    )
    first = month_label(months[0]).split()[0]
    last = month_label(months[-1]).split()[0]
    period_name = f"{first}–{last} 2026"
    ws = wb.active
    ws.title = "Yearly Summary"
    write_vendor_sheet(
        ws,
        f"{client_name} — Vendor Spend Yearly Summary ({period_name})",
        f"Year Total Purchases: {year_total:,.2f} | Months: {', '.join(month_label(m) for m in months)}",
        year_rows,
        year_total,
        "% of Year Total",
        extra_col="Months Active",
    )
    for mm in months:
        rows = monthly[mm]
        mtotal = sum(r["total"] for r in rows)
        mws = wb.create_sheet(month_label(mm)[:31])
        write_vendor_sheet(
            mws,
            f"{client_name} — {month_label(mm)} Vendor Spend",
            f"Month Total Purchases: {mtotal:,.2f}",
            rows,
            mtotal,
            "% of Month Total",
        )
    out_path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(out_path)
    return out_path


def build_all(root: Path):
    created = []
    for vp in sorted(root.rglob("Vendor Purchases")):
        pdfs = sorted(vp.glob("*.pdf"))
        if not pdfs:
            continue
        client_folder = vp.parent.name
        by_month = defaultdict(list)
        for pdf in pdfs:
            m = re.search(r"(\d{2})2026\.pdf$", pdf.name)
            if m:
                by_month[m.group(1)].append(pdf)
        if not by_month:
            continue
        slug = re.sub(r"[^a-z0-9]+", "", client_folder.lower())[:24] or "client"
        monthly_data = {}
        for mm, plist in sorted(by_month.items()):
            vendors_map = defaultdict(lambda: {"orders": 0, "total": 0.0})
            for pdf in plist:
                parsed, _ = parse_pdf_tables(pdf)
                for v in parsed:
                    vendors_map[v["vendor"]]["orders"] += v["orders"]
                    vendors_map[v["vendor"]]["total"] += v["total"]
            monthly_data[mm] = sorted(
                [
                    {"vendor": n, "orders": d["orders"], "total": d["total"]}
                    for n, d in vendors_map.items()
                ],
                key=lambda x: -x["total"],
            )
        # remove prior vendor_spend outputs for this client
        for old in vp.glob("*vendor_spend*.xlsx"):
            old.unlink()
        out = vp / f"{slug}_2026_vendor_spend.xlsx"
        build_client_workbook(client_folder, monthly_data, out)
        created.append(out)
    return created


if __name__ == "__main__":
    root = Path(__file__).resolve().parents[1] / "reports" / "Clients"
    created = build_all(root)
    print("FILES", len(created))
    for p in created:
        print(p)
