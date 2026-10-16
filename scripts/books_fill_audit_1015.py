#!/usr/bin/env python3
"""Fill books audit gaps (2026-10-16 run): open-month days + known holes.

Uses S2K Daily PDFs + expanded DLY under --pdf-dir / --dly-dir.
Sales fill-once unless --force-day marks a zero placeholder row for refill.
"""
from __future__ import annotations

import argparse
import json
import shutil
import sys
from datetime import date, datetime
from pathlib import Path

import openpyxl

REPO = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(REPO / "scripts"))
from fill_daily_excel_from_softsp import (  # noqa: E402
    STORES,
    apply_invoices,
    dly_stamp,
    find_dly,
    load_deduct_from_workbook,
    load_snapshot,
    parse_bd_store,
    parse_dly_invoices,
    parse_single,
    pdf_stamp,
    purch_for_excel,
    rules_to_json,
    scan_excel_dir,
    write_values,
    header_map,
    HDR,
    template_formulas,
)
MONTH_NAMES = {
    1: "January",
    2: "February",
    3: "March",
    4: "April",
    5: "May",
    6: "June",
    7: "July",
    8: "August",
    9: "September",
    10: "October",
    11: "November",
    12: "December",
}


def month_sheet(year: int, month: int) -> str:
    return f"{MONTH_NAMES[month]} {year}"


def month_source_sheet(year: int, month: int) -> str:
    return f"{MONTH_NAMES[month]} {year} Source"


def month_calc_sheet(year: int, month: int) -> str:
    return f"{MONTH_NAMES[month]} Calculations"


def row_has_sales(ws, r: int, sales_col: str = "E") -> bool:
    gas = ws[f"B{r}"].value
    sales = ws[f"{sales_col}{r}"].value
    if sales not in (None, "", 0):
        return True
    if gas not in (None, "", 0):
        return True
    return False


def fill_brookhurst_month(
    wb,
    d: date,
    vals: dict,
    purch: float | None,
    *,
    update_purch: bool,
    force: bool,
) -> list[str]:
    src_name = month_source_sheet(d.year, d.month)
    main_name = month_sheet(d.year, d.month)
    src = wb[src_name]
    ws = wb[main_name]
    day = d.day
    sr = 2 + day
    mr = HDR + day
    changed = []
    if force or not row_has_sales(src, sr):
        if src[f"A{sr}"].value in (None, ""):
            src[f"A{sr}"] = datetime(d.year, d.month, d.day)
            changed.append(f"src!A{sr}=date")
        tax_sum = sum(
            (vals.get(k) or 0.0) for k in ("tax1", "tax4", "scratch", "lotto", "card")
        )
        mapping = {
            "B": vals.get("gas_vol"),
            "C": vals.get("gas_profit"),
            "D": (
                round(vals["gas_profit"] / vals["gas_vol"], 4)
                if vals.get("gas_vol") and vals.get("gas_profit") is not None
                else None
            ),
            "E": (
                round(vals["cstore_total"] - tax_sum, 2)
                if vals.get("cstore_total") is not None
                else None
            ),
            "J": vals.get("cstore_total"),
            "K": vals.get("tax1", 0.0),
            "L": vals.get("tax4", 0.0),
            "M": vals.get("scratch", 0.0),
            "N": vals.get("lotto", 0.0),
            "O": vals.get("card", 0.0),
        }
        for col, val in mapping.items():
            if val is None:
                continue
            cur = src[f"{col}{sr}"].value
            if force or cur in (None, "", 0):
                src[f"{col}{sr}"] = val
                changed.append(f"src!{col}{sr}={val}")
    if update_purch or purch is not None:
        shown = purch_for_excel(purch)
        src[f"F{sr}"] = shown
        ws[f"F{mr}"] = f"='{src_name}'!F{sr}"
        changed.append(f"purch F sr={sr}")
    for col, formula in {
        "G": f'=IF(COUNTA(A{sr}:F{sr})=0,"",E{sr}-F{sr})',
        "H": f'=IFERROR(G{sr}/E{sr},"")',
        "I": f'=IF(COUNTA(A{sr}:F{sr})=0,"",C{sr}+G{sr})',
    }.items():
        if src[f"{col}{sr}"].value in (None, ""):
            src[f"{col}{sr}"] = formula
    links = {
        "B": f"='{src_name}'!B{sr}",
        "C": f"='{src_name}'!C{sr}",
        "E": f"='{src_name}'!E{sr}",
        "J": f"='{src_name}'!J{sr}",
        "K": f"='{src_name}'!K{sr}",
        "L": f"='{src_name}'!L{sr}",
        "M": f"='{src_name}'!M{sr}",
        "N": f"='{src_name}'!N{sr}",
        "O": f"='{src_name}'!O{sr}",
    }
    for col, formula in links.items():
        if ws[f"{col}{mr}"].value in (None, ""):
            ws[f"{col}{mr}"] = formula
    for col, formula in template_formulas(ws, day).items():
        if col in ("D", "G", "H", "I") and ws[f"{col}{mr}"].value in (None, ""):
            ws[f"{col}{mr}"] = formula
    return changed


def repair_brookhurst_from_main(wb, d: date) -> list[str]:
    """When Source row is corrupt but main sheet has sane values, sync Source from main."""
    src_name = month_source_sheet(d.year, d.month)
    main_name = month_sheet(d.year, d.month)
    if src_name not in wb.sheetnames or main_name not in wb.sheetnames:
        return []
    src = wb[src_name]
    main = wb[main_name]
    sr = 2 + d.day
    mr = HDR + d.day
    ms = main[f"E{mr}"].value
    ss = src[f"E{sr}"].value
    if ms in (None, "", 0):
        return []
    try:
        if ss is not None and float(ss) == float(ms):
            return []
        if ss is not None and float(ss) < float(ms) * 5:
            return []
    except (TypeError, ValueError):
        pass
    changed = []
    for col in ("B", "C", "E", "F", "J", "K", "L", "M", "N", "O"):
        v = main[f"{col}{mr}"].value
        if v is None:
            continue
        if isinstance(v, str) and v.startswith("="):
            continue
        src[f"{col}{sr}"] = v
        changed.append(f"repair src!{col}{sr}<-main")
    return changed


def fill_san_diego_month(
    wb, d: date, vals: dict, purch: float | None, *, update_purch: bool, force: bool
) -> list[str]:
    sn = month_calc_sheet(d.year, d.month)
    ws = wb[sn]
    r = 2 + d.day
    if not force and row_has_sales(ws, r):
        if update_purch or purch is not None:
            ws[f"F{r}"] = purch_for_excel(purch)
            return [f"F{r}=purch"]
        return []
    changed = []
    tax_sum = sum(
        (vals.get(k) or 0.0) for k in ("tax1", "tax4", "scratch", "lotto", "card")
    )
    mapping = [
        ("B", vals.get("gas_vol")),
        ("C", vals.get("gas_profit")),
        ("J", vals.get("cstore_total")),
        ("K", vals.get("tax1", 0.0)),
        ("L", vals.get("tax4", 0.0)),
        ("M", vals.get("scratch", 0.0)),
        ("N", vals.get("lotto", 0.0)),
        ("O", vals.get("card", 0.0)),
    ]
    if vals.get("cstore_total") is not None:
        mapping.append(("E", round(vals["cstore_total"] - tax_sum, 2)))
    for col, val in mapping:
        if val is None:
            continue
        if ws[f"{col}{r}"].value in (None, "", 0) or force:
            ws[f"{col}{r}"] = val
            changed.append(f"{col}{r}={val}")
    if update_purch or purch is not None:
        ws[f"F{r}"] = purch_for_excel(purch)
        changed.append(f"F{r}=purch")
    return changed


def fill_standard_month(
    wb,
    sid: str,
    d: date,
    vals: dict,
    purch: float | None,
    *,
    update_purch: bool,
    force: bool,
) -> list[str]:
    sn = month_sheet(d.year, d.month)
    ws = wb[sn]
    cols = header_map(ws)
    r = HDR + d.day
    if not force and row_has_sales(ws, r, cols.get("gas_vol", "B") and "B" or "B"):
        gas_col = cols.get("gas_vol", "B")
        sales_col = cols.get("sales") or "E"
        gas = ws[f"{gas_col}{r}"].value
        sales = ws[f"{sales_col}{r}"].value
        if (gas not in (None, "", 0)) or (sales not in (None, "", 0)):
            if update_purch or purch is not None:
                pc = cols.get("purch") or "F"
                ws[f"{pc}{r}"] = purch_for_excel(purch)
                return [f"{pc}{r}=purch"]
            return []
    return write_values(
        ws, d.day, vals, cols, purch=purch, update_purch=update_purch, force=force
    )


def parse_days(raw: str) -> list[date]:
    return [date.fromisoformat(p.strip()) for p in raw.split(",") if p.strip()]


def main() -> None:
    ap = argparse.ArgumentParser()
    ap.add_argument("--xlsx-dir", required=True, type=Path)
    ap.add_argument("--pdf-dir", required=True, type=Path)
    ap.add_argument("--dly-dir", type=Path)
    ap.add_argument("--days", required=True, help="Comma ISO dates")
    ap.add_argument("--force-days", default="", help="ISO dates to refill even if row exists")
    ap.add_argument("--out-dir", type=Path)
    ap.add_argument("--only")
    ap.add_argument("--repair-brookhurst-main", action="store_true")
    ap.add_argument("--report", type=Path, default=Path("/tmp/books_work/fill_report.json"))
    args = ap.parse_args()
    days = parse_days(args.days)
    force_days = set(parse_days(args.force_days)) if args.force_days else set()
    out_dir = args.out_dir or args.xlsx_dir
    out_dir.mkdir(parents=True, exist_ok=True)
    only = {s.strip() for s in args.only.split(",")} if args.only else None
    through = max(days)
    rules = load_snapshot()
    purch_by_store: dict[str, dict[str, float | None]] = {}
    if args.dly_dir:
        bd_inv = None
        for sid, meta in STORES.items():
            if only and sid not in only:
                continue
            _name, _fname, _mode, _dsub, _tso, dly_sub = meta
            rule = rules.get(sid) or load_deduct_from_workbook(
                args.xlsx_dir / _fname, store_id=sid
            )
            if dly_sub == "BD_central":
                if bd_inv is None:
                    p = find_dly(args.dly_dir, "BD_central", through)
                    bd_inv = parse_dly_invoices(p) if p else []
                invs = [
                    i
                    for i in bd_inv
                    if i["store_id"] == sid or i["store_id"].startswith(sid)
                ]
            else:
                p = find_dly(args.dly_dir, dly_sub, through)
                invs = parse_dly_invoices(p, {sid}) if p else []
            purch_by_store[sid] = apply_invoices(rule, invs)

    report = []
    for sid, (name, fname, mode, sub, tso, _dly_sub) in STORES.items():
        if only and sid not in only:
            continue
        src_path = args.xlsx_dir / fname
        if not src_path.exists():
            report.append({"id": sid, "error": "missing_xlsx"})
            continue
        dest = out_dir / fname
        if dest.resolve() != src_path.resolve():
            shutil.copy2(src_path, dest)
        wb = openpyxl.load_workbook(dest)
        store_log = {"id": sid, "filled": [], "repairs": []}
        store_purch = purch_by_store.get(sid, {})
        do_purch = sid in purch_by_store

        if args.repair_brookhurst_main and sid == "42098":
            for d in days:
                ch = repair_brookhurst_from_main(wb, d)
                if ch:
                    store_log["repairs"].append({"date": d.isoformat(), "changed": ch})

        for d in days:
            force = d in force_days
            pdf = args.pdf_dir / sub / pdf_stamp(d)
            purch = store_purch.get(d.isoformat())
            vals: dict = {}
            src_pdf = pdf if pdf.exists() else None
            if not src_pdf:
                bd_pdf = args.pdf_dir / "BD_central" / pdf_stamp(d)
                if bd_pdf.exists():
                    src_pdf = bd_pdf
            if src_pdf:
                if mode == "single":
                    # Never parse_single on BD_central — that PDF is multi-store totals.
                    if "BD_central" in src_pdf.parts:
                        vals = parse_bd_store(src_pdf, sid) or {}
                    else:
                        vals = parse_single(src_pdf)
                else:
                    vals = parse_bd_store(src_pdf, tso or sid) or {}
            elif force:
                store_log["filled"].append({"date": d.isoformat(), "error": "no_pdf"})
                if do_purch:
                    if sid == "42048":
                        sn = month_calc_sheet(d.year, d.month)
                        ws = wb[sn]
                        ws[f"F{2 + d.day}"] = purch_for_excel(purch)
                    elif sid == "42098":
                        src = wb[month_source_sheet(d.year, d.month)]
                        src[f"F{2 + d.day}"] = purch_for_excel(purch)
                continue
            else:
                if do_purch and sid not in ("42048", "42098"):
                    ws = wb[month_sheet(d.year, d.month)]
                    cols = header_map(ws)
                    pc = cols.get("purch") or "F"
                    ws[f"{pc}{HDR + d.day}"] = purch_for_excel(purch)
                continue

            if not vals.get("cstore_total") and not force:
                store_log["filled"].append({"date": d.isoformat(), "error": "parse_miss"})
                continue

            if sid == "42048":
                ch = fill_san_diego_month(
                    wb, d, vals, purch, update_purch=do_purch, force=force
                )
            elif sid == "42098":
                ch = fill_brookhurst_month(
                    wb, d, vals, purch, update_purch=do_purch, force=force
                )
            else:
                ch = fill_standard_month(
                    wb, sid, d, vals, purch, update_purch=do_purch, force=force
                )
            if ch:
                store_log["filled"].append({"date": d.isoformat(), "changed": ch})

        if store_log["filled"] or store_log["repairs"]:
            wb.save(dest)
        report.append(store_log)

    args.report.parent.mkdir(parents=True, exist_ok=True)
    args.report.write_text(json.dumps(report, indent=2) + "\n")
    print(f"wrote {args.report}")


if __name__ == "__main__":
    main()
