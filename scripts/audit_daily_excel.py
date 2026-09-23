#!/usr/bin/env python3
"""Audit Daily Excel workbooks vs a books sync payload for one target day.

Reads C-Store Sales (E, or component formulas J−K−…) and Net Daily Purchases (F).
Formula-only F cells are reported as mismatches unless --write-purch is set, which
overwrites F with the payload purch for that day (Excel SoT alignment).

Usage:
  python3 scripts/audit_daily_excel.py \\
    --xlsx-dir /tmp/books_excel_0921 \\
    --payload /tmp/s2k/exports/sync_excel_plus_0921.json \\
    --day 2026-09-21

  python3 scripts/audit_daily_excel.py ... --day 2026-09-21 --write-purch
"""
from __future__ import annotations

import argparse
import json
import re
import sys
from datetime import date, datetime
from pathlib import Path

try:
    import openpyxl
    from openpyxl.worksheet.formula import ArrayFormula
except ImportError as exc:  # pragma: no cover
    raise SystemExit("openpyxl required: pip install openpyxl") from exc

sys.path.insert(0, str(Path(__file__).resolve().parent))
from sync_from_excel import STORE_FILES, find_daily_header, find_month_sheet, parse_date

REPO = Path(__file__).resolve().parents[1]


def cell_val(v):
    if isinstance(v, ArrayFormula):
        return getattr(v, "text", str(v))
    return v


def is_formula(v) -> bool:
    v = cell_val(v)
    return isinstance(v, str) and v.startswith("=")


def as_num(v):
    v = cell_val(v)
    if isinstance(v, (int, float)) and not isinstance(v, bool):
        return float(v)
    return None


def parse_d(raw):
    d = parse_date(raw)
    if d:
        return d
    if isinstance(raw, str):
        m = re.search(r"DATE\(\s*(\d+)\s*,\s*(\d+)\s*,\s*(\d+)", raw, re.I)
        if m:
            return date(int(m.group(1)), int(m.group(2)), int(m.group(3)))
    return None


def resolve_sheet(wb, sid: str, month: str):
    if sid == "42048":
        calc = None
        for name in wb.sheetnames:
            if "september" in name.lower() and "calc" in name.lower():
                calc = name
                break
        if calc:
            ws = wb[calc]
            hr, cols = find_daily_header(ws)
            return ws, hr or 2, cols or {"date": 1, "sales": 5, "purch": 6}
    sn = find_month_sheet(wb, month)
    if not sn:
        raise SystemExit(f"{sid}: no sheet for {month}")
    ws = wb[sn]
    hr, cols = find_daily_header(ws)
    if not hr:
        return ws, 8, {"date": 1, "sales": 5, "purch": 6}
    return ws, hr, cols


def find_row(ws, hr, cols, target: date):
    for r in range(hr + 1, hr + 45):
        d = parse_d(cell_val(ws.cell(r, cols["date"]).value))
        if d == target:
            return r
    return None


def sales_resolve(ws, r, sid: str):
    """Resolve C-Store Sales from numeric E or component columns."""
    scol = 5
    sv = cell_val(ws.cell(r, scol).value)
    n = as_num(sv)
    if n is not None:
        return n, "numeric"
    # Spring Mtn: I − J − K (I = C-Store Total)
    if sid == "42280":
        i, j, k = as_num(ws.cell(r, 9).value), as_num(ws.cell(r, 10).value), as_num(
            ws.cell(r, 11).value
        )
        if i is not None:
            return round(i - (j or 0) - (k or 0), 2), "I-J-K"
    # Vista: J − K − L
    if sid == "42438":
        j, k, l = as_num(ws.cell(r, 10).value), as_num(ws.cell(r, 11).value), as_num(
            ws.cell(r, 12).value
        )
        if j is not None:
            return round(j - (k or 0) - (l or 0), 2), "J-K-L"
    # Default: J − K − L − M − N − O
    vals = [as_num(ws.cell(r, c).value) for c in range(10, 16)]
    if vals[0] is not None:
        return round(vals[0] - sum(v or 0 for v in vals[1:]), 2), "J-O"
    # Spring-style if I looks like total
    i = as_num(ws.cell(r, 9).value)
    if i is not None and is_formula(sv) and "I" in str(sv):
        j, k = as_num(ws.cell(r, 10).value), as_num(ws.cell(r, 11).value)
        return round(i - (j or 0) - (k or 0), 2), "I-J-K"
    return None, "unresolved"


def find_xlsx(xlsx_dir: Path, sid: str) -> Path:
    cands = list(xlsx_dir.glob(f"{sid}_daily.xlsx")) + list(xlsx_dir.glob(f"*{sid}*Daily*.xlsx"))
    if not cands:
        raise SystemExit(f"no workbook for {sid} in {xlsx_dir}")
    return cands[0]


def main() -> None:
    ap = argparse.ArgumentParser(description=__doc__)
    ap.add_argument("--xlsx-dir", required=True, type=Path)
    ap.add_argument("--payload", required=True, type=Path, help="sync JSON with stations[].days")
    ap.add_argument("--day", required=True, help="YYYY-MM-DD")
    ap.add_argument("--write-purch", action="store_true", help="overwrite F with payload purch")
    ap.add_argument(
        "--out",
        type=Path,
        default=REPO / "scripts/exports/excel_day_audit.json",
    )
    args = ap.parse_args()
    target = date.fromisoformat(args.day)
    month = target.strftime("%Y-%m")
    payload = json.loads(args.payload.read_text())
    site_by = {s["id"]: s for s in payload["stations"]}

    rows = []
    ok_s = ok_p = 0
    writes = []

    print(
        f"{'ID':6} {'Store':22} {'ExSales':>10} {'SiSales':>10} {'S':4} "
        f"{'ExPurch':>10} {'SiPurch':>10} {'P':4}"
    )
    print("-" * 90)

    for sid, (name, _) in STORE_FILES.items():
        path = find_xlsx(args.xlsx_dir, sid)
        wb = openpyxl.load_workbook(path)
        ws, hr, cols = resolve_sheet(wb, sid, month)
        r = find_row(ws, hr, cols, target)
        site = site_by.get(sid)
        site_day = None
        if site:
            site_day = next((d for d in site["days"] if d["date"] == args.day), None)
        if r is None or site_day is None:
            rows.append(
                {
                    "id": sid,
                    "name": name,
                    "err": "no row" if r is None else "no site day",
                }
            )
            print(f"{sid:6} {name:22} ERR {rows[-1]['err']}")
            continue

        sales, via = sales_resolve(ws, r, sid)
        pcol = cols.get("purch", 6)
        old_p = cell_val(ws.cell(r, pcol).value)
        purch = as_num(old_p)

        if args.write_purch:
            want = float(site_day["purch"])
            if purch is None or abs(purch - want) > 0.02 or is_formula(old_p):
                ws.cell(r, pcol).value = want
                # San Diego display sheet mirror
                if sid == "42048":
                    sn = find_month_sheet(wb, month)
                    if sn:
                        dws = wb[sn]
                        dws.cell(8 + target.day, 6).value = want
                # Brookhurst: fix broken Source!E3 sales link on target day
                if sid == "42098" and is_formula(ws.cell(r, 5).value):
                    raw_e = str(cell_val(ws.cell(r, 5).value) or "")
                    if "Source'!E3" in raw_e or sales is None:
                        ws.cell(r, 5).value = f"=J{r}-K{r}-L{r}-M{r}-N{r}-O{r}"
                        sales, via = sales_resolve(ws, r, sid)
                purch = want
                writes.append(f"{sid} F{r}={want}")
                wb.save(path)

        ms = sales is not None and abs(sales - site_day["sales"]) < 0.05
        mp = purch is not None and abs(purch - site_day["purch"]) < 0.05
        if ms:
            ok_s += 1
        if mp:
            ok_p += 1
        rows.append(
            {
                "id": sid,
                "name": name,
                "row": r,
                "excel_sales": sales,
                "site_sales": site_day["sales"],
                "excel_purch": purch,
                "site_purch": site_day["purch"],
                "match_sales": ms,
                "match_purch": mp,
                "sales_via": via,
                "purch_was_formula": is_formula(old_p),
            }
        )
        print(
            f"{sid:6} {name:22} "
            f"{(sales if sales is not None else float('nan')):10.2f} "
            f"{site_day['sales']:10.2f} {'OK' if ms else 'DIFF':4} "
            f"{(purch if purch is not None else float('nan')):10.2f} "
            f"{site_day['purch']:10.2f} {'OK' if mp else 'DIFF':4}"
        )

    summary = {
        "day": args.day,
        "sales_ok": ok_s,
        "purch_ok": ok_p,
        "total": len(STORE_FILES),
        "writes": writes,
    }
    print(f"\nSales OK: {ok_s}/{len(STORE_FILES)}  Purch OK: {ok_p}/{len(STORE_FILES)}")
    if writes:
        print(f"Wrote purch on {len(writes)} cells")

    args.out.parent.mkdir(parents=True, exist_ok=True)
    args.out.write_text(json.dumps({"summary": summary, "rows": rows}, indent=2) + "\n")
    print(f"wrote {args.out}")
    if ok_s < len(STORE_FILES) or ok_p < len(STORE_FILES):
        sys.exit(1)


if __name__ == "__main__":
    main()
