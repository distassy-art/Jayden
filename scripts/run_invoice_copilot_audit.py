#!/usr/bin/env python3
"""Stage Daily Book Summary + DLY + scans for Copilot, then OCR-append Audit.xlsx."""
from __future__ import annotations

import argparse
import json
import os
import re
import ssl
import subprocess
import sys
import tempfile
import urllib.error
import urllib.parse
import urllib.request
from datetime import date, datetime
from pathlib import Path
from zoneinfo import ZoneInfo

import openpyxl

ROOT = Path(__file__).resolve().parents[1]
CONFIG_PATH = ROOT / "scripts" / "invoice_scan_clients.json"
COOKIES_PATH = Path(os.environ.get("OD_COOKIES", "/tmp/od_cookies.json"))
HOST = "https://smartsolutionsai26-my.sharepoint.com"
BASE = f"{HOST}/personal/minamorcos_smartsolutionsai26_onmicrosoft_com"
DOCS = "/personal/minamorcos_smartsolutionsai26_onmicrosoft_com/Documents"
COPILOT_ROOT = f"{DOCS}/Microsoft Copilot Chat Files"  # flat files; Copilot folder rejects subfolders
PT = ZoneInfo("America/Los_Angeles")
SSL_CTX = ssl.create_default_context()

VENDOR_CANON = {
    "coca cola": "REYES COCA COLA BOTTLING, LLC",
    "coke": "REYES COCA COLA BOTTLING, LLC",
    "reyes coca": "REYES COCA COLA BOTTLING, LLC",
    "pepsi": "PEPSICO BEVERAGE SALES LLC",
    "pepsico": "PEPSICO BEVERAGE SALES LLC",
    "frito": "FRITO-LAY",
    "frito-lay": "FRITO-LAY",
    "frito lay": "FRITO-LAY",
    "bimbo": "BIMBO BAKERIES USA, INC",
    "harbor": "HARBOR DISTRIBUTING, LLC",
    "kadi": "KADI",
    "northstar": "NORTHSTAR",
    "north star": "NORTHSTAR",
    "united express": "UNITED EXPRESS",
    "ab one": "AB SALES",
    "ab sales": "AB SALES",
    "red bull": "RED BULL DISTRIBUTION COMPANY INC",
    "7up": "7UP",
    "7-up": "7UP",
    "universal": "UNIVERSAL",
    "bower": "BOWER",
    "singtech": "BOWER",
}


def load_config() -> dict:
    return json.loads(CONFIG_PATH.read_text())


def cookie_header() -> str:
    cookies = json.loads(COOKIES_PATH.read_text())
    return "; ".join(
        f"{c['name']}={c['value']}"
        for c in cookies
        if "sharepoint.com" in c.get("domain", "")
    )


def http(url: str, method: str = "GET", data: bytes | None = None, headers: dict | None = None, timeout: int = 180):
    h = {
        "Cookie": cookie_header(),
        "Accept": "application/json;odata=verbose",
        "User-Agent": "Mozilla/5.0",
    }
    if headers:
        h.update(headers)
    req = urllib.request.Request(url, data=data, headers=h, method=method)
    try:
        with urllib.request.urlopen(req, context=SSL_CTX, timeout=timeout) as resp:
            return resp.status, resp.read()
    except urllib.error.HTTPError as e:
        return e.code, e.read()


def form_digest() -> str:
    code, body = http(BASE + "/_api/contextinfo", method="POST", data=b"")
    if code != 200:
        raise RuntimeError(f"digest {code} {body[:200]}")
    return json.loads(body)["d"]["GetContextWebInformation"]["FormDigestValue"]


def list_files(folder_rel: str) -> list[dict]:
    url = (
        BASE
        + f"/_api/web/GetFolderByServerRelativeUrl('{urllib.parse.quote(folder_rel, safe='/')}')"
        + "/Files?$select=Name,ServerRelativeUrl,Length,TimeLastModified&$top=500"
    )
    code, body = http(url)
    if code == 404:
        return []
    if code != 200:
        raise RuntimeError(f"list {code} {folder_rel} {body[:160]}")
    return json.loads(body)["d"]["results"]


def ensure_folder(folder_rel: str) -> None:
    parts = folder_rel[len(DOCS):].strip("/").split("/")
    cur = DOCS
    dig = form_digest()
    for part in parts:
        parent = cur
        cur = f"{cur}/{part}"
        code, _ = http(
            BASE
            + f"/_api/web/GetFolderByServerRelativeUrl('{urllib.parse.quote(cur, safe='/')}')"
            + "?$select=Name"
        )
        if code == 200:
            continue
        add = (
            BASE
            + f"/_api/web/GetFolderByServerRelativeUrl('{urllib.parse.quote(parent, safe='/')}')"
            + "/Folders/add('"
            + urllib.parse.quote(part)
            + "')"
        )
        code, body = http(add, method="POST", data=b"", headers={"X-RequestDigest": dig})
        if code not in (200, 201):
            code2, _ = http(
                BASE
                + f"/_api/web/GetFolderByServerRelativeUrl('{urllib.parse.quote(cur, safe='/')}')"
                + "?$select=Name"
            )
            if code2 != 200:
                raise RuntimeError(f"mkdir {cur} {code} {body[:200]}")


def download_bytes(server_rel: str) -> bytes:
    url = (
        BASE
        + f"/_api/web/GetFileByServerRelativeUrl('{urllib.parse.quote(server_rel, safe='/')}')/$value"
    )
    code, body = http(url, timeout=300)
    if code != 200:
        raise RuntimeError(f"download {code} {server_rel} {body[:160]}")
    return body


def upload_bytes(folder_rel: str, name: str, data: bytes) -> None:
    ensure_folder(folder_rel)
    dig = form_digest()
    url = (
        BASE
        + f"/_api/web/GetFolderByServerRelativeUrl('{urllib.parse.quote(folder_rel, safe='/')}')"
        + f"/Files/add(url='{urllib.parse.quote(name)}',overwrite=true)"
    )
    code, body = http(
        url,
        method="POST",
        data=data,
        headers={"X-RequestDigest": dig, "Content-Type": "application/octet-stream"},
        timeout=300,
    )
    if code not in (200, 201):
        raise RuntimeError(f"upload {code} {folder_rel}/{name} {body[:200]}")


def ocr_pdf(path: Path, max_pages: int = 2) -> str:
    with tempfile.TemporaryDirectory() as td:
        prefix = Path(td) / "p"
        subprocess.run(
            ["pdftoppm", "-png", "-r", "200", "-f", "1", "-l", str(max_pages), str(path), str(prefix)],
            check=False,
            capture_output=True,
        )
        texts = []
        for img in sorted(Path(td).glob("p*.png")):
            r = subprocess.run(
                ["tesseract", str(img), "stdout", "--psm", "6"],
                capture_output=True,
                text=True,
            )
            texts.append(r.stdout or "")
        return "\n".join(texts)


def parse_amount(text: str) -> float | None:
    labeled = re.findall(
        r"(?:amount\s*due(?:\s*for\s*this\s*invoice)?|invoice\s*total|total\s*due|grand\s*total|balance\s*due|\btotal\b)\s*[:\$]?\s*\$?\s*([0-9]{1,3}(?:,[0-9]{3})*(?:\.[0-9]{2}))",
        text,
        flags=re.I,
    )
    if labeled:
        try:
            return float(labeled[-1].replace(",", ""))
        except ValueError:
            pass
    amounts = [float(a.replace(",", "")) for a in re.findall(r"\$\s*([0-9]{1,3}(?:,[0-9]{3})*\.[0-9]{2})", text)]
    amounts = [v for v in amounts if 1 <= v <= 50000]
    return max(amounts) if amounts else None


def parse_invoice_date(text: str) -> date | None:
    months = {
        "jan": 1, "feb": 2, "mar": 3, "apr": 4, "may": 5, "jun": 6,
        "jul": 7, "aug": 8, "sep": 9, "oct": 10, "nov": 11, "dec": 12,
    }
    patterns = [
        r"(?:invoice\s*date|delivery\s*date|date)\s*[:\s]*(\d{1,2})[/-](\d{1,2})[/-](\d{2,4})",
        r"\b(\d{1,2})[/-](\d{1,2})[/-](20\d{2})\b",
        r"\b(\d{1,2})\s+(Jan|Feb|Mar|Apr|May|Jun|Jul|Aug|Sep|Oct|Nov|Dec)[a-z]*\s+(20\d{2})\b",
        r"\b(Jan|Feb|Mar|Apr|May|Jun|Jul|Aug|Sep|Oct|Nov|Dec)[a-z]*\s+(\d{1,2}),?\s+(20\d{2})\b",
    ]
    for pat in patterns:
        m = re.search(pat, text, flags=re.I)
        if not m:
            continue
        g = m.groups()
        try:
            if g[1].isalpha():
                d, mon, y = int(g[0]), months[g[1][:3].lower()], int(g[2])
            elif g[0].isalpha():
                mon, d, y = months[g[0][:3].lower()], int(g[1]), int(g[2])
            else:
                a, b, y = int(g[0]), int(g[1]), int(g[2])
                if y < 100:
                    y += 2000
                mon, d = (b, a) if a > 12 else (a, b)
            return date(y, mon, d)
        except Exception:
            continue
    return None


def canon_vendor(text_or_name: str, aliases: dict[str, str]) -> str:
    low = (text_or_name or "").lower()
    for k, v in sorted(VENDOR_CANON.items(), key=lambda kv: -len(kv[0])):
        if k in low:
            return v
    for k, v in sorted(aliases.items(), key=lambda kv: -len(kv[0])):
        if k.lower() in low:
            return v
    return (text_or_name or "UNKNOWN").strip().upper()[:60]


def client_paths(client: dict, month: str) -> dict:
    dest = client["dest"]
    return {
        "scans": f"{DOCS}/{dest}/Scans/{month}",
        "daily_summary": f"{DOCS}/{dest}/{month}/Daily Summary",
        "audit_xlsx": f"{DOCS}/{dest}/{client['name']} Audit.xlsx",
        "dly_bd": f"{DOCS}/Clients/BIG DADDY/PDF/dly/{month}",
        "daily_bd": f"{DOCS}/Clients/BIG DADDY/PDF/daily/{month}",
    }


def pick_files(paths: dict, as_of: date) -> dict:
    mmddyyyy = f"{as_of.month:02d}{as_of.day:02d}{as_of.year:04d}"
    ddmm = f"{as_of.day:02d}{as_of.month:02d}"
    out = {"daily": [], "dly": [], "scans": []}

    for folder in (paths["daily_summary"], paths["daily_bd"]):
        try:
            files = list_files(folder)
        except Exception:
            files = []
        for f in files:
            n = f["Name"].lower()
            if "dly" in n or "dpt" in n:
                continue
            if n.startswith(mmddyyyy.lower()) and n.endswith(".pdf"):
                out["daily"].append(f)
        if out["daily"]:
            break

    for folder in (paths["daily_summary"], paths["dly_bd"]):
        try:
            files = list_files(folder)
        except Exception:
            files = []
        for f in files:
            n = f["Name"].lower().replace("-", "").replace("_", "")
            if "dly" in n and mmddyyyy.lower() in n:
                out["dly"].append(f)
        if out["dly"]:
            break

    try:
        scans = list_files(paths["scans"])
    except Exception:
        scans = []
    # Include renamed scans for as_of and prior 3 calendar days (invoice date DDMM),
    # plus raw Scan* for as_of.
    from datetime import timedelta
    ddmm_set = set()
    for i in range(0, 4):
        d = as_of - timedelta(days=i)
        ddmm_set.add(f"{d.day:02d}{d.month:02d}")
    for f in scans:
        n = f["Name"]
        if not n.lower().endswith(".pdf"):
            continue
        if n.lower().startswith("scan"):
            if re.match(rf"Scan{as_of.year:04d}-{as_of.month:02d}-{as_of.day:02d}", n, re.I):
                out["scans"].append(f)
            continue
        if any(re.search(rf"\b{dd}(?:-\d+)?\.pdf$", n, re.I) for dd in ddmm_set):
            out["scans"].append(f)
        elif mmddyyyy in n.replace("-", ""):
            out["scans"].append(f)
    return out


def stage_for_copilot(client: dict, as_of: date, picked: dict) -> dict:
    """Stage into Microsoft Copilot Chat Files as flat prefixed files.

    Copilot Chat Files does not reliably support nested folders for upload,
    so we use: Audit_{ClientSanitized}_{as-of}_01_daily_book_summary.pdf etc.
    """
    safe_client = re.sub(r"[^A-Za-z0-9]+", "", client["name"]) or "Client"
    prefix = f"Audit_{safe_client}_{as_of.isoformat()}_"
    uploaded = []
    for i, f in enumerate(picked["daily"], 1):
        name = f"{prefix}01_daily_book_summary.pdf" if i == 1 else f"{prefix}01_daily_book_summary_{i}.pdf"
        upload_bytes(COPILOT_ROOT, name, download_bytes(f["ServerRelativeUrl"]))
        uploaded.append(name)
    for i, f in enumerate(picked["dly"], 1):
        name = f"{prefix}02_dly.pdf" if i == 1 else f"{prefix}02_dly_{i}.pdf"
        upload_bytes(COPILOT_ROOT, name, download_bytes(f["ServerRelativeUrl"]))
        uploaded.append(name)
    for i, f in enumerate(picked["scans"], 1):
        safe = re.sub(r"[^\w.\-]+", "_", f["Name"]).strip("_")
        name = f"{prefix}03_scan_{i:02d}_{safe}"
        upload_bytes(COPILOT_ROOT, name, download_bytes(f["ServerRelativeUrl"]))
        uploaded.append(name)
    prompt = (
        f"Client invoice audit — {client['name']} — as of {as_of.isoformat()}\n\n"
        f"Upload these Copilot Chat Files (already staged with prefix {prefix}):\n"
        f"1) {prefix}01_daily_book_summary.pdf — Daily Book Summary\n"
        f"2) {prefix}02_dly.pdf — DLY (None Fuel Invoice Total)\n"
        f"3) {prefix}03_scan_*.pdf — renamed scanned invoices\n\n"
        "Task:\n"
        "- From each scanned invoice extract Vendor Name, Invoice Date, Amount.\n"
        f"- Append new rows into SCANNED INVOICES on the {as_of.strftime('%B %Y')} sheet of \"{client['name']} Audit.xlsx\".\n"
        "- From the DLY, reconcile active non-fuel S2K vendor amounts "
        "(Matched / Missing in S2K / Amount mismatch; $5 tolerance).\n"
        f"- Update Cover / month note with through-date {as_of.isoformat()}.\n"
        "- Do not overwrite older matched scanned rows; only add new invoices.\n"
    )
    upload_bytes(COPILOT_ROOT, f"{prefix}COPILOT_AUDIT_PROMPT.txt", prompt.encode("utf-8"))
    uploaded.append(f"{prefix}COPILOT_AUDIT_PROMPT.txt")
    return {"stage_folder": COPILOT_ROOT, "prefix": prefix, "uploaded": uploaded}



def extract_scan_rows(picked: dict, aliases: dict[str, str]) -> list[dict]:
    rows = []
    with tempfile.TemporaryDirectory() as td:
        td_path = Path(td)
        for f in picked["scans"]:
            local = td_path / f["Name"]
            local.write_bytes(download_bytes(f["ServerRelativeUrl"]))
            text = ocr_pdf(local)
            stem = Path(f["Name"]).stem
            vendor_guess = re.sub(r"\s+\d{4}(?:-\d+)?$", "", stem).strip()
            if vendor_guess.lower().startswith("scan"):
                vendor_guess = ""
            vendor = canon_vendor(text if not vendor_guess else vendor_guess + "\n" + text, aliases)
            inv_date = parse_invoice_date(text)
            m = re.search(r"(\d{2})(\d{2})(?:-\d+)?$", stem)
            if not inv_date and m and not stem.lower().startswith("scan"):
                d, mon = int(m.group(1)), int(m.group(2))
                try:
                    inv_date = date(datetime.now(PT).year, mon, d)
                except ValueError:
                    pass
            amount = parse_amount(text)
            rows.append(
                {
                    "vendor": vendor,
                    "invoice_date": inv_date,
                    "amount": amount,
                    "source": f["Name"],
                    "ocr_sample": " ".join(text.split())[:220],
                }
            )
    return rows


def append_audit_rows(client: dict, month: str, rows: list[dict], as_of: date, dry_run: bool = False) -> dict:
    paths = client_paths(client, month)
    candidates = [paths["audit_xlsx"], f"{DOCS}/{client['dest']}/{client['name']} Audit.xlsx"]
    data = used = None
    for rel in candidates:
        try:
            data = download_bytes(rel)
            used = rel
            break
        except Exception:
            continue
    if data is None:
        return {"status": "error", "error": "audit workbook not found", "tried": candidates}

    local = Path(f"/tmp/audit_{client.get('station','x')}.xlsx")
    local.write_bytes(data)
    wb = openpyxl.load_workbook(local)
    sheet_name = f"{as_of.strftime('%B')} {as_of.year}"
    if sheet_name not in wb.sheetnames:
        return {"status": "error", "error": f"missing sheet {sheet_name}", "sheets": wb.sheetnames}
    ws = wb[sheet_name]

    header_row = None
    for r in range(1, 40):
        val = ws.cell(r, 1).value
        if val and str(val).strip().upper() == "VENDOR NAME":
            header_row = r
            break
    if header_row is None:
        return {"status": "error", "error": "VENDOR NAME header not found"}

    existing = set()
    write_row = header_row + 1
    for r in range(header_row + 1, header_row + 200):
        v = ws.cell(r, 1).value
        if v is None or str(v).strip() == "" or str(v).strip().upper().startswith("MTD"):
            write_row = r
            break
        d = ws.cell(r, 2).value
        a = ws.cell(r, 3).value
        if isinstance(d, datetime):
            dkey = d.date().isoformat()
        elif isinstance(d, date):
            dkey = d.isoformat()
        else:
            dkey = str(d)
        try:
            aval = round(float(a), 2) if a is not None else None
        except Exception:
            aval = None
        existing.add((str(v).strip().upper(), dkey, aval))
        write_row = r + 1

    added = []
    for row in rows:
        if row["amount"] is None or row["invoice_date"] is None:
            continue
        key = (str(row["vendor"]).strip().upper(), row["invoice_date"].isoformat(), round(float(row["amount"]), 2))
        if key in existing:
            continue
        if not dry_run:
            ws.cell(write_row, 1).value = row["vendor"]
            ws.cell(write_row, 2).value = datetime.combine(row["invoice_date"], datetime.min.time())
            ws.cell(write_row, 3).value = float(row["amount"])
        added.append({**row, "row": write_row, "status": "dry_run" if dry_run else "added"})
        existing.add(key)
        write_row += 1

    ws.cell(3, 1).value = (
        f"{as_of.strftime('%B')} scanned invoices loaded through {as_of.isoformat()}; "
        f"agent OCR append {datetime.now(PT).strftime('%Y-%m-%d %H:%M %Z')}."
    )
    if "Cover" in wb.sheetnames:
        cover = wb["Cover"]
        for r in range(1, 20):
            v = cover.cell(r, 1).value
            if v and as_of.strftime("%b") in str(v):
                cover.cell(r, 2).value = (
                    f"{as_of.strftime('%B')} scanned invoices loaded through {as_of.isoformat()}."
                )
                break

    if not dry_run and added:
        out_path = Path(f"/tmp/audit_{client.get('station','x')}_out.xlsx")
        wb.save(out_path)
        folder, name = used.rsplit("/", 1)
        upload_bytes(folder, name, out_path.read_bytes())

    return {
        "status": "ok",
        "audit_file": used,
        "added": [
            {
                "vendor": a["vendor"],
                "invoice_date": a["invoice_date"].isoformat() if a.get("invoice_date") else None,
                "amount": a.get("amount"),
                "source": a.get("source"),
                "row": a.get("row"),
                "status": a.get("status"),
            }
            for a in added
        ],
        "incomplete": [
            {
                "source": r["source"],
                "vendor": r["vendor"],
                "amount": r["amount"],
                "invoice_date": r["invoice_date"].isoformat() if r["invoice_date"] else None,
            }
            for r in rows
            if r["amount"] is None or r["invoice_date"] is None
        ],
    }


def run(only: str | None, as_of_s: str | None, month: str | None, dry_run: bool) -> int:
    cfg = load_config()
    aliases = {k.lower(): v for k, v in (cfg.get("vendor_aliases") or {}).items()}
    clients = list(cfg.get("clients") or [])
    if only:
        clients = [c for c in clients if only.lower() in c["name"].lower() or only == str(c.get("station") or "")]
    as_of = date.fromisoformat(as_of_s) if as_of_s else datetime.now(PT).date()
    month = month or f"{as_of.year:04d}-{as_of.month:02d} {as_of.strftime('%B')}"

    report = {"started": datetime.now(PT).isoformat(), "as_of": as_of.isoformat(), "month": month, "clients": []}
    for client in clients:
        paths = client_paths(client, month)
        picked = pick_files(paths, as_of)
        entry = {
            "client": client["name"],
            "picked": {
                "daily": [f["Name"] for f in picked["daily"]],
                "dly": [f["Name"] for f in picked["dly"]],
                "scans": [f["Name"] for f in picked["scans"]],
            },
        }
        if dry_run:
            entry["stage"] = {"status": "dry_run"}
        else:
            entry["stage"] = stage_for_copilot(client, as_of, picked)
        rows = extract_scan_rows(picked, aliases) if picked["scans"] else []
        entry["extract"] = [
            {
                "vendor": r["vendor"],
                "invoice_date": r["invoice_date"].isoformat() if r["invoice_date"] else None,
                "amount": r["amount"],
                "source": r["source"],
                "ocr_sample": r.get("ocr_sample"),
            }
            for r in rows
        ]
        entry["audit"] = append_audit_rows(client, month, rows, as_of, dry_run=dry_run)
        report["clients"].append(entry)
        print(json.dumps(entry, indent=2, default=str, ensure_ascii=False))

    Path("/tmp/invoice_copilot_audit_report.json").write_text(json.dumps(report, indent=2, default=str, ensure_ascii=False))
    print("Wrote /tmp/invoice_copilot_audit_report.json")
    return 0


def main() -> None:
    ap = argparse.ArgumentParser(description=__doc__)
    ap.add_argument("--only")
    ap.add_argument("--as-of", help="YYYY-MM-DD")
    ap.add_argument("--month", help='e.g. "2026-09 September"')
    ap.add_argument("--dry-run", action="store_true")
    args = ap.parse_args()
    sys.exit(run(args.only, args.as_of, args.month, args.dry_run))


if __name__ == "__main__":
    main()
