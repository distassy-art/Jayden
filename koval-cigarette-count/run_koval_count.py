#!/usr/bin/env python3
"""Koval (Flamingo #42279) cigarette count 2026-09-28 — S2K post + Jayden packet."""
from __future__ import annotations

import json
import shutil
import sys
import urllib.parse
from datetime import date, datetime
from decimal import Decimal, ROUND_HALF_UP
from pathlib import Path

from fpdf import FPDF
from openpyxl import load_workbook

sys.path.insert(0, "/tmp/cig-inv")
import build_all  # noqa: E402

WORKDIR = Path("/tmp/cig-inv/koval-count")
OUT = Path("/workspace/koval-cigarette-count/42279/2026-09-28")
STATION = "42279"
NAME = "Flamingo (Koval)"
SITEID = 195
ACCOUNT = -121
TRANREF = "CIG COUNT 09282026"
TRANDATE = "2026-09-28"
REPORT_END = "2026-09-28"

# Grid quantities from scan (PDF page order: p2 rows 34–69, p1 rows 70–105, p3 rows 106–107).
# Rows 1–33 are not in the scanned PDF (first sheet page missing); treat as blank except Turquoise.
# Row 57: handwritten "?" — excluded (ambiguous).
HANDWRITTEN: dict[int, int | None] = {
    8: 8,  # American Spirit Turquoise — store policy; only below-grid line posted from extras
    34: 8,
    35: 4,
    36: 5,
    37: 8,
    38: 28,
    39: 0,
    40: 10,
    41: 9,
    42: 4,
    43: 23,
    44: 4,
    45: 1,
    46: 11,
    47: 9,
    48: 11,
    49: 15,
    50: 10,
    51: 7,
    52: 10,
    53: 38,
    54: 9,
    55: 6,
    56: 7,
    58: 26,
    59: 49,
    60: 20,
    61: 0,
    62: 0,
    63: 10,
    64: 13,
    65: 9,
    66: 35,
    67: 11,
    68: 7,
    69: 11,
    70: 3,
    71: 19,
    72: 0,
    73: 0,
    74: 6,
    75: 2,
    76: 0,
    77: 0,
    78: 0,
    79: 7,
    80: 5,
    81: 5,
    82: 13,
    83: 50,
    84: 51,
    85: 11,
    86: 7,
    87: 3,
    88: 5,
    89: 5,
    90: 3,
    91: 7,
    92: 0,
    93: 9,
    94: 5,
    95: 8,
    96: 8,
    97: 7,
    98: 5,
    99: 14,
    100: 5,
    101: 6,
    102: 1,
    103: 0,
    104: 0,
    105: 11,
    106: 6,
    107: 8,
}

AMBIGUOUS_ROWS = [57]
BELOW_GRID_NOT_POSTED = [
    "Handwritten list below grid on scan page 3 (Camel, American Spirit, Crown, Kool, etc.) "
    "was not posted except American Spirit Turquoise on row 8 per store policy.",
    "Scan PDF is missing the printed grid for rows 1–33; those lines were left unchanged in S2K.",
]


def money(n: float) -> str:
    sign = "-" if n < 0 else ""
    return f"{sign}${abs(n):,.2f}"


def qty(n) -> str:
    if isinstance(n, float) and n.is_integer():
        return str(int(n))
    return str(n)


def round2(x: float) -> float:
    return float(Decimal(str(x)).quantize(Decimal("0.01"), rounding=ROUND_HALF_UP))


def load_sheet_items() -> list[dict]:
    rows = []
    for line in (WORKDIR / "sheet-items.tsv").read_text().splitlines():
        n, upc, desc = line.split("\t", 2)
        rows.append({"n": int(n), "upc": upc.strip(), "desc": desc.strip()})
    return rows


def pull_packs_xlsx(session) -> Path:
    build_all.END = REPORT_END
    depts = build_all.cig_depts(session)
    pack_dept = next((d for d, n in depts if "PACK" in n.upper()), depts[0][0])
    sid = build_all.session_id(session)
    url = (
        "https://store.s2kprime.com/report?rpt="
        + urllib.parse.quote("SKU Sales By Item")
        + "&id="
        + urllib.parse.quote(str(sid))
        + f"&StartDate=2026-08-01&EndDate={REPORT_END}&StationID={SITEID}&Dept={pack_dept}&Toggle=1&rs:Format=EXCELOPENXML"
    )
    r = session.get(url, timeout=240)
    if r.content[:2] != b"PK":
        raise RuntimeError(f"packs report failed {r.status_code}")
    path = WORKDIR / "packs.xlsx"
    path.write_bytes(r.content)
    return path


def parse_qoh(path: Path) -> dict[str, dict]:
    wb = load_workbook(path, data_only=True)
    ws = wb.active
    latest: dict[str, dict] = {}
    for row in ws.iter_rows(min_row=13, values_only=True):
        if not row or len(row) < 9:
            continue
        desc, sku, dept, sdate, qoh, sold, cost = (
            row[1],
            row[3],
            row[4],
            row[5],
            row[6],
            row[7],
            row[8],
        )
        if not isinstance(sku, str) or not sku.strip().isdigit():
            continue
        if not isinstance(sdate, datetime):
            continue
        upc = sku.strip()
        rec = {
            "desc": (desc or "").strip(),
            "dept": (dept or "").strip(),
            "dt": sdate.date().isoformat(),
            "qoh": float(qoh or 0),
            "cost": float(cost or 0),
        }
        prev = latest.get(upc)
        if not prev or rec["dt"] >= prev["dt"]:
            latest[upc] = rec
    return latest


def product_lookup(session, upc: str) -> dict | None:
    r = session.get(
        "https://store.s2kprime.com/api/product",
        params={"upc": upc},
        timeout=40,
    )
    if r.status_code != 200:
        return None
    data = r.json()
    if not isinstance(data, list) or not data:
        return None
    for hit in data:
        if str(hit.get("upc") or "") == upc:
            return hit
    return data[0]


def find_existing_tran(session) -> dict | None:
    xsrf = session.cookies.get("XSRF-TOKEN")
    hdr = {"X-XSRF-TOKEN": xsrf} if xsrf else {}
    for url in (
        f"https://store.s2kprime.com/api/tran/item-10?siteid={SITEID}&fromdate={TRANDATE}&todate={TRANDATE}",
        f"https://store.s2kprime.com/api/tran/item?siteid={SITEID}&trandate={TRANDATE}",
    ):
        r = session.get(url, headers=hdr, timeout=60)
        if r.status_code != 200:
            continue
        try:
            data = r.json()
        except json.JSONDecodeError:
            continue
        items = data if isinstance(data, list) else data.get("data") or []
        for it in items:
            if str(it.get("tranref") or "") == TRANREF and int(it.get("siteid") or 0) == SITEID:
                return it
    return None


def post_inventory(session, lines: list[dict]) -> dict:
    meta_path = WORKDIR / "post-meta.json"
    if meta_path.exists():
        try:
            prior = json.loads(meta_path.read_text())
            if (
                prior.get("tranref") == TRANREF
                and int(prior.get("siteid") or 0) == SITEID
                and prior.get("id")
            ):
                return {
                    "skipped": True,
                    "id": prior["id"],
                    "reason": "already posted (local post-meta)",
                }
        except (json.JSONDecodeError, TypeError, ValueError):
            pass
    existing = find_existing_tran(session)
    if existing:
        return {"skipped": True, "id": existing.get("id"), "reason": "already posted"}

    tran_lines = []
    for row in lines:
        tran_lines.append(
            {
                "varid": row["varid"],
                "qty": float(row["actual"]),
                "cal_type": 0,
                "props": json.dumps({"sold": 0}),
                "model_dirty": True,
            }
        )
    item = {
        "siteid": SITEID,
        "tranref": TRANREF,
        "trandate": TRANDATE,
        "trandate2": TRANDATE,
        "trantype": 10,
        "lines": tran_lines,
        "deleted_lines": [],
        "props": json.dumps(
            {
                "note": (
                    f"{NAME} cigarette physical count written {TRANDATE}. "
                    "Quantity on hand set to the counted packs. "
                    "American Spirit Turquoise included. "
                    "Blank quantity cells were left unchanged. "
                    "Rows 1–33 grid page missing from scan."
                )
            }
        ),
        "model_dirty": True,
    }
    body = {
        "type": "item",
        "trantype": 10,
        "data": [item],
        "create_new": True,
        "return_get": True,
    }
    xsrf = session.cookies.get("XSRF-TOKEN")
    hdr = {"X-XSRF-TOKEN": xsrf} if xsrf else {}
    for path in (
        "/api/tran/update/item-10-undefined",
        "/api/tran/update",
    ):
        r = session.post(
            "https://store.s2kprime.com" + path,
            json=body,
            headers=hdr,
            timeout=120,
        )
        if r.status_code == 200 and "error" not in r.text.lower():
            data = r.json()
            if isinstance(data, list) and data:
                return {"skipped": False, "id": data[0].get("id"), "response": data[0]}
            if isinstance(data, dict) and data.get("id"):
                return {"skipped": False, "id": data.get("id"), "response": data}
        if "duplicate" in r.text.lower():
            return {"skipped": True, "reason": "duplicate", "text": r.text[:400]}
    raise RuntimeError(f"post failed last status {r.status_code} {r.text[:400]}")


class Sheet(FPDF):
    def footer(self):
        self.set_y(-12)
        self.set_font("Helvetica", "", 8)
        self.set_x(self.l_margin)
        self.cell(self.epw, 8, f"{self.page_no()}/{{nb}}", align="R")


def write_pdfs(lines: list[dict], summary: dict, tranid: int | None):
    cols = [12, 86, 16, 16, 14, 20, 26]
    headers = ["#", "Description", "Actual", "Book", "Diff", "Unit", "Amount"]

    def table(pdf, rows, amount_key):
        pdf.set_font("Helvetica", "B", 8)
        pdf.set_x(pdf.l_margin)
        for w, label in zip(cols, headers):
            pdf.cell(w, 5, label, new_x="RIGHT", new_y="TOP")
        pdf.ln(5)
        pdf.set_font("Helvetica", "", 8)
        for row in rows:
            if pdf.get_y() > 262:
                pdf.add_page()
            vals = [
                str(row["n"]),
                row["desc"][:42],
                qty(row["actual"]),
                qty(row["qoh"]),
                qty(row["diff"]),
                f"{row['cost']:.4f}",
                money(row[amount_key]),
            ]
            pdf.set_x(pdf.l_margin)
            for w, val in zip(cols, vals):
                pdf.cell(w, 4.4, val, new_x="RIGHT", new_y="TOP")
            pdf.ln(4.4)

    pdf = Sheet(format="Letter", unit="mm")
    pdf.alias_nb_pages()
    pdf.set_auto_page_break(auto=True, margin=14)
    pdf.set_margins(10, 12, 10)
    pdf.add_page()
    pdf.set_font("Helvetica", "B", 14)
    pdf.cell(0, 8, f"{NAME} #{STATION} cigarette count", new_x="LMARGIN", new_y="NEXT")
    pdf.set_font("Helvetica", "", 10)
    pdf.multi_cell(
        0,
        5,
        f"September 28, 2026. Actual {money(summary['actualCost'])} for {summary['actualPacks']} packs. "
        f"Overall {money(summary['overall'])}.",
    )
    actual_rows = sorted(lines, key=lambda r: r["n"])
    for row in actual_rows:
        row["actual_amt"] = round2(row["actual"] * row["cost"])
    pdf.ln(2)
    pdf.set_font("Helvetica", "B", 11)
    pdf.cell(0, 6, "Actual cost", new_x="LMARGIN", new_y="NEXT")
    table(pdf, actual_rows, "actual_amt")
    minus_rows = sorted([r for r in lines if r["cost_delta"] < 0], key=lambda r: r["cost_delta"])
    plus_rows = sorted([r for r in lines if r["cost_delta"] > 0], key=lambda r: -r["cost_delta"])
    pdf.ln(2)
    pdf.set_font("Helvetica", "B", 11)
    pdf.cell(0, 6, "Minus, book above the count", new_x="LMARGIN", new_y="NEXT")
    table(pdf, minus_rows, "cost_delta")
    pdf.ln(2)
    pdf.set_font("Helvetica", "B", 11)
    pdf.cell(0, 6, "Plus, count above the book", new_x="LMARGIN", new_y="NEXT")
    table(pdf, plus_rows, "cost_delta")
    pdf.output(OUT / "count.pdf")

    pdf2 = Sheet(format="Letter", unit="mm")
    pdf2.alias_nb_pages()
    pdf2.set_auto_page_break(auto=True, margin=14)
    pdf2.set_margins(10, 12, 10)
    pdf2.add_page()
    pdf2.set_font("Helvetica", "B", 14)
    pdf2.cell(0, 8, "SKU Inventory", new_x="LMARGIN", new_y="NEXT")
    pdf2.set_font("Helvetica", "B", 12)
    pdf2.cell(0, 6, f"{NAME} #{STATION}", new_x="LMARGIN", new_y="NEXT")
    pdf2.set_font("Helvetica", "", 10)
    pdf2.cell(0, 5, f"September 28, 2026. Reference {TRANREF}.", new_x="LMARGIN", new_y="NEXT")
    if tranid:
        pdf2.cell(0, 5, f"Transaction {tranid}.", new_x="LMARGIN", new_y="NEXT")
    cols2 = [12, 100, 36, 36]
    pdf2.ln(2)
    pdf2.set_font("Helvetica", "B", 8)
    for w, label in zip(cols2, ["#", "Description", "UPC", "Quantity on hand"]):
        pdf2.cell(w, 5, label, new_x="RIGHT", new_y="TOP")
    pdf2.ln(5)
    pdf2.set_font("Helvetica", "", 8)
    for row in sorted(lines, key=lambda r: r["n"]):
        if pdf2.get_y() > 262:
            pdf2.add_page()
        pdf2.set_x(pdf2.l_margin)
        for w, val in zip(
            cols2,
            [str(row["n"]), row["desc"][:48], row["upc"], qty(row["actual"])],
        ):
            pdf2.cell(w, 4.4, val, new_x="RIGHT", new_y="TOP")
        pdf2.ln(4.4)
    pdf2.output(OUT / "s2k-change.pdf")


def build_entry(summary: dict, tranid: int | None) -> dict:
    base = f"/inventory/{STATION}/{TRANDATE}"
    return {
        "stationId": STATION,
        "client": NAME,
        "date": TRANDATE,
        "dateLabel": "September 28, 2026",
        "category": "Cigarettes",
        "heading": (
            f"Count written September 28, 2026. SKU Inventory {TRANREF}"
            + (f", transaction {tranid}." if tranid else ".")
            + " American Spirit Turquoise included. Rows 1–33 missing from scan were unchanged. "
            "Below-grid counts were not posted except Turquoise."
        ),
        "summary": summary["narrative"],
        "count": {
            "actualPacks": summary["actualPacks"],
            "actualCost": summary["actualCost"],
            "minusPacks": summary["minusPacks"],
            "minusCost": summary["minusCost"],
            "plusPacks": summary["plusPacks"],
            "plusCost": summary["plusCost"],
            "overall": summary["overall"],
        },
        "files": [
            {
                "group": "Scans",
                "label": "Scan page 1 (grid rows 70–105)",
                "href": f"{base}/scan-1.pdf",
            },
            {
                "group": "Scans",
                "label": "Scan page 2 (grid rows 34–69)",
                "href": f"{base}/scan-2.pdf",
            },
            {
                "group": "Scans",
                "label": "Scan page 3 (rows 106–107 and below-grid notes)",
                "href": f"{base}/scan-3.pdf",
            },
            {
                "group": "S2K quantity on hand",
                "label": "Packs, department 200",
                "href": f"{base}/qoh-packs.pdf",
            },
            {
                "group": "S2K quantity on hand",
                "label": "Cartons, department 281",
                "href": f"{base}/qoh-cartons.pdf",
            },
            {
                "group": "S2K quantity on hand",
                "label": "Quantity on hand change",
                "href": f"{base}/s2k-change.pdf",
            },
            {
                "group": "Count",
                "label": "Cigarette count",
                "href": f"{base}/count.pdf",
            },
        ],
    }


def main():
    import argparse

    ap = argparse.ArgumentParser()
    ap.add_argument("--no-post", action="store_true", help="Build packet only; do not call S2K tran/update")
    args = ap.parse_args()
    OUT.mkdir(parents=True, exist_ok=True)
    items = load_sheet_items()
    logins = build_all.load_logins()
    session, _ = build_all.login(*logins["hotmail"])
    build_all.switch(session, ACCOUNT)

    packs_path = pull_packs_xlsx(session)
    qoh_by_upc = parse_qoh(packs_path)
    (WORKDIR / "latest-qoh.json").write_text(json.dumps(qoh_by_upc, indent=2))

    lines: list[dict] = []
    skipped_no_varid = []
    for it in items:
        n = it["n"]
        actual = HANDWRITTEN.get(n)
        if actual is None:
            continue
        prod = product_lookup(session, it["upc"])
        if not prod:
            skipped_no_varid.append(n)
            continue
        varid = prod["id"]
        q = qoh_by_upc.get(it["upc"], {})
        qoh = float(q.get("qoh", 0))
        cost = float(q.get("cost") or 0)
        if cost <= 0 and prod.get("cost"):
            cost = float(prod["cost"])
        diff = float(actual) - qoh
        lines.append(
            {
                "n": n,
                "upc": it["upc"],
                "desc": it["desc"],
                "actual": actual,
                "qoh": qoh,
                "diff": diff,
                "cost": cost,
                "cost_delta": round2(diff * cost),
                "varid": varid,
                "qoh_asof": q.get("dt"),
                "dept": q.get("dept", "CIGARETTES - PACKS"),
            }
        )

    actual_packs = sum(r["actual"] for r in lines)
    actual_cost = round2(sum(r["actual"] * r["cost"] for r in lines))
    minus = [r for r in lines if r["diff"] < 0]
    plus = [r for r in lines if r["diff"] > 0]
    minus_packs = sum(-r["diff"] for r in minus)
    plus_packs = sum(r["diff"] for r in plus)
    minus_cost = round2(sum(-r["cost_delta"] for r in minus if r["cost_delta"] < 0))
    plus_cost = round2(sum(r["cost_delta"] for r in plus if r["cost_delta"] > 0))
    overall = round2(plus_cost - minus_cost)

    if args.no_post:
        meta_path = WORKDIR / "post-meta.json"
        tranid = None
        post = {"skipped": True, "reason": "--no-post"}
        if meta_path.exists():
            try:
                tranid = json.loads(meta_path.read_text()).get("id")
            except json.JSONDecodeError:
                pass
    else:
        post = post_inventory(session, lines)
        tranid = post.get("id")

    summary = {
        "station": STATION,
        "name": NAME,
        "siteid": SITEID,
        "date": TRANDATE,
        "tranref": TRANREF,
        "tranid": tranid,
        "linesPosted": len(lines),
        "blankLeftUnchanged": 107 - len(lines),
        "actualPacks": int(actual_packs),
        "actualCost": actual_cost,
        "minusPacks": int(minus_packs),
        "minusCost": minus_cost,
        "plusPacks": int(plus_packs),
        "plusCost": plus_cost,
        "overall": overall,
        "turquoise": "included, count 8 from below-grid (row 8 grid page missing)",
        "ambiguousRows": AMBIGUOUS_ROWS,
        "missingGridRows": "1-33",
        "post": post,
        "narrative": (
            f"Actual {money(actual_cost)} for {int(actual_packs)} packs on {len(lines)} lines. "
            f"Minus {money(minus_cost)} ({int(minus_packs)} packs), plus {money(plus_cost)} "
            f"({int(plus_packs)} packs). Overall {money(overall)}."
        ),
    }
    (WORKDIR / "summary.json").write_text(json.dumps(summary, indent=2))
    (OUT / "lines.json").write_text(json.dumps(lines, indent=2))
    (WORKDIR / "post-meta.json").write_text(
        json.dumps(
            {
                "id": tranid,
                "tranref": TRANREF,
                "trandate": TRANDATE,
                "siteid": SITEID,
                "nlines": len(lines),
                "status": 200 if tranid else None,
                "post": post,
            },
            indent=2,
        )
    )

    write_pdfs(lines, summary, tranid)

    sid = build_all.session_id(session)
    for dept_id, label, fname in (
        ("200", "packs", "qoh-packs.pdf"),
        ("281", "cartons", "qoh-cartons.pdf"),
    ):
        url = (
            "https://store.s2kprime.com/report?rpt="
            + urllib.parse.quote("SKU Sales By Item")
            + "&id="
            + urllib.parse.quote(str(sid))
            + f"&StartDate=2026-08-01&EndDate={REPORT_END}&StationID={SITEID}&Dept={dept_id}&Toggle=1&rs:Format=PDF"
        )
        rp = session.get(url, timeout=240)
        if rp.content[:4] == b"%PDF":
            (OUT / fname).write_bytes(rp.content)
    (OUT / "entry.json").write_text(json.dumps(build_entry(summary, tranid), indent=2) + "\n")

    # Scans: PDF page 1 = grid 70-105, page 2 = 34-69, page 3 = 106-107 + notes
    shutil.copy(WORKDIR / "Koval_Cigarette_Count_09282026.pdf", OUT / "count-sheet.pdf")
    for i, src in enumerate(
        [
            WORKDIR / "pdf-page-1.png",
            WORKDIR / "pdf-page-2.png",
            WORKDIR / "pdf-page-3.png",
        ],
        1,
    ):
        pdf_scan = OUT / f"scan-{i}.pdf"
        from PIL import Image

        im = Image.open(src).convert("RGB")
        tmp = OUT / f"_scan{i}.png"
        im.save(tmp)
        fp = FPDF(unit="mm", format="Letter")
        fp.add_page()
        fp.image(str(tmp), x=0, y=0, w=215.9)
        fp.output(str(pdf_scan))
        tmp.unlink(missing_ok=True)

    readme = Path("/workspace/koval-cigarette-count/README.md")
    readme.write_text(
        f"""# {NAME} cigarette count — September 28, 2026

Station {STATION}. S2K site {SITEID}.

- Reference: `{TRANREF}`
- Transaction: {tranid or '(see post-meta.json)'}

## Totals

- Actual: {summary['actualPacks']} packs, {money(summary['actualCost'])}
- Minus: {summary['minusPacks']} packs, {money(summary['minusCost'])}
- Plus: {summary['plusPacks']} packs, {money(summary['plusCost'])}
- Overall: {money(summary['overall'])}
- Lines posted: {summary['linesPosted']}
- Rows 1–33: printed grid not in scan PDF — QoH left unchanged ({33 - 1} lines) except American Spirit Turquoise row 8 (count 8).
- Row 57 ({items[56]['desc']}): ambiguous handwritten quantity — not posted.
- {BELOW_GRID_NOT_POSTED[0]}

## Files

Packet directory: `42279/2026-09-28/`.

Billing and site-publish are out of scope for this branch.
"""
    )

    print(json.dumps(summary, indent=2))
    if skipped_no_varid:
        print("no varid for rows", skipped_no_varid, file=sys.stderr)


if __name__ == "__main__":
    main()
